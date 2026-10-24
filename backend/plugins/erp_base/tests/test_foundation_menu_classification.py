from plugins.erp_base.source_services import stage
"""分类菜单、共表保存、原值承接及缺表迁移在独立库验证。"""
from io import BytesIO
from django.apps import apps
from django.test import TestCase
from django.db import connection
from django.core.management import call_command
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient
from openpyxl import load_workbook
from system.models import Menu, Role, Dept, RoleMenuPermission, RoleMenuButtonPermission
from system.excel import write_workbook
from plugins.erp_base.foundation_catalog import DEFINITIONS, CATALOG
from plugins.erp_base.foundation_initialize import load
from plugins.erp_base.source_excel_history import columns
from plugins.erp_base.legacy_models import LegacyImportBatch
from plugins.erp_base.source_context import SourceDocumentContext
from plugins.erp_base.models import BusinessAction
from .test_phase_report_audit import create as source_fixture

def profile(table, **filters):
    return next((x for x in DEFINITIONS.values() if x.get('source_table') == table and x['filters'] == filters))

class FoundationClassificationTests(TestCase):

    @classmethod
    def setUpTestData(cls):
        call_command('init_data', verbosity=0)
        cls.actor = get_user_model().objects.get(username='superadmin')

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.actor)

    def url(self, item, suffix=''):
        return '/api/system/' + item['entity'] + '/' + suffix + '?_foundation=' + item['key']

    def ok(self, response):
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        return response.data['data']

    def count(self, response):
        self.assertEqual(response.status_code, 200, response.data)
        return response.data['total']

    def create(self, item, values):
        return self.client.post(self.url(item), values, format='json')

    def test_catalogue_covers_all_sources_and_original_fields_without_extra_legacy_id(self):
        self.assertEqual(len(CATALOG['plan_sources']), 158)
        self.assertEqual(len(DEFINITIONS), 499)
        for item in DEFINITIONS.values():
            m = apps.get_model(item['model'])
            self.assertTrue(set(item['filters']) <= set(m.SOURCE_FIELDS.values()), item['key'])
            self.assertNotIn('legacy_id', {f.name for f in m._meta.fields})
            self.assertNotIn('旧erpID', {str(f.verbose_name) for f in m._meta.fields})

    def test_all_source_tables_and_all_definition_menus_exist_after_migration(self):
        tables = set(connection.introspection.table_names())
        for m in apps.get_models():
            if hasattr(m, 'SOURCE_TABLE'):
                self.assertIn(m._meta.db_table, tables, m._meta.label)
        for key, item in DEFINITIONS.items():
            menu = Menu.objects.get(component_name=key)
            self.assertEqual(menu.resource, item['entity'])
            self.assertEqual(menu.component, 'plugins/erp/FoundationEntry')

    def test_menu_rebuild_is_idempotent_and_preserves_old_urls_and_business_rows(self):
        unit = profile('dye.WidthGramWeight')
        m = apps.get_model(unit['model'])
        m.objects.create(old_code='dup', old_name_c='同名', old_kind=None)
        m.objects.create(old_code='dup', old_name_c='同名', old_kind='')
        paths = dict(Menu.objects.values_list('component_name', 'web_path'))
        total = Menu.objects.count()
        load()
        self.assertEqual(Menu.objects.count(), total)
        self.assertEqual(dict(Menu.objects.values_list('component_name', 'web_path')), paths)
        self.assertEqual(list(m.objects.order_by('pk').values_list('old_kind', flat=True)), [None, ''])

    def test_all_499_business_definitions_and_lists_are_accessible(self):
        for key, item in DEFINITIONS.items():
            response = self.client.get('/api/system/' + item['entity'] + '/foundation_entry/', {'key': key})
            self.assertEqual(self.ok(response)['key'], key)
            response = self.client.get(self.url(item), {'limit': 1})
            self.assertEqual(response.status_code, 200, (key, response.data))

    def test_business_shortcuts_share_the_same_rows(self):
        shortcut = DEFINITIONS['foundation_short_purchase_supplier']
        original = DEFINITIONS[shortcut['shortcut_of']]
        row = self.ok(self.create(shortcut, {'old_num': 'SUP-2', 'old_short_name': '快捷入口供应商'}))
        self.assertEqual(self.client.get(self.url(original)).data['data'][0]['id'], row['id'])
        self.assertEqual(apps.get_model(original['model']).objects.count(), 1)

    def test_wrong_entity_or_unknown_definition_is_explicitly_rejected(self):
        unit = profile('dye.WidthGramWeight')
        response = self.client.get('/api/system/' + unit['entity'] + '/?_foundation=not-registered')
        self.assertEqual(response.status_code, 400)
        other = profile('dye.p_Customer')
        response = self.client.get('/api/system/' + unit['entity'] + '/?_foundation=' + other['key'])
        self.assertEqual(response.status_code, 400)

    def test_width_and_weight_are_separate_and_keep_duplicate_names_and_codes(self):
        width = profile('dye.WidthGramWeight', old_kind='幅宽')
        weight = profile('dye.WidthGramWeight', old_kind='克重')
        self.ok(self.create(width, {'old_code': '004', 'old_name_c': 'CM'}))
        self.ok(self.create(width, {'old_code': '004', 'old_name_c': 'CM'}))
        self.ok(self.create(weight, {'old_code': '001', 'old_name_c': 'G回潮后勿超', 'old_num_value': None}))
        m = apps.get_model(width['model'])
        self.assertEqual(m.objects.count(), 3)
        self.assertEqual(m.objects.filter(old_kind='幅宽').count(), 2)
        self.assertEqual(m.objects.get(old_kind='克重').old_name_c, 'G回潮后勿超')
        self.assertEqual(self.count(self.client.get(self.url(width))), 2)
        self.assertEqual(self.count(self.client.get(self.url(weight))), 1)

    def test_wrong_kind_cannot_be_silently_overwritten_or_updated_through_other_scope(self):
        width = profile('dye.WidthGramWeight', old_kind='幅宽')
        weight = profile('dye.WidthGramWeight', old_kind='克重')
        response = self.create(width, {'old_code': '001', 'old_kind': '克重'})
        self.assertEqual(response.status_code, 400)
        row = self.ok(self.create(weight, {'old_code': '001', 'old_name_c': 'G'}))
        response = self.client.patch(self.url(width, str(row['id']) + '/'), {'old_name_c': '不能跨类'}, format='json')
        self.assertEqual(response.status_code, 404)
        self.assertEqual(apps.get_model(width['model']).objects.get().old_name_c, 'G')

    def test_dictionary_scope_preserves_old_case_and_trailing_spaces(self):
        item = profile('dye.p_FieldValid', old_field_type='Transport')
        self.ok(self.create(item, {'old_field_type': 'transport  ', 'old_code': 'D', 'old_name': '运输'}))
        row = apps.get_model(item['model']).objects.get()
        self.assertEqual(row.old_field_type, 'transport  ')
        self.assertEqual(self.count(self.client.get(self.url(item))), 1)

    def test_packaging_dimensions_and_repeated_names_remain_independent(self):
        item = profile('dye.p_FieldValid', old_field_type='PackBag')
        for code, weight in [('A10', '1.1'), ('A12', '1.2')]:
            self.ok(self.create(item, {'old_code': code, 'old_name': '孟加拉', 'old_num_value': weight, 'old_long_cm': '60', 'old_wide_cm': '45', 'old_high_cm': '24'}))
        rows = list(apps.get_model(item['model']).objects.order_by('old_code'))
        self.assertEqual([r.old_name for r in rows], ['孟加拉', '孟加拉'])
        self.assertNotEqual(rows[0].old_num_value, rows[1].old_num_value)
        self.assertEqual(rows[0].old_long_cm, 60)

    def test_route_parent_and_original_step_relation_are_kept(self):
        item = profile('dye.p_FieldValid', old_field_type='Production')
        parent = self.ok(self.create(item, {'old_code': 'PT01', 'old_name': '装缸→染色'}))
        self.assertTrue(BusinessAction.objects.filter(action='entry_route_name_pending').exists())
        source_fixture('p_Process', ProcessID='GX01', Name='染色')
        m = apps.get_model('erp_base.SourceDyePFieldValidSub')
        response = self.client.post('/api/system/erp_base_sourcedyepfieldvalidsub/', {'_entry': 'p_ProductionDetail', 'source_parent_1': parent['id'], 'old_sub_code': 'GX01', 'old_sort_value': 1}, format='json')
        self.ok(response)
        child = m.objects.get()
        self.assertEqual((child.old_field_type, child.old_code, child.old_sub_code), ('Production', 'PT01', 'GX01'))
        self.assertEqual(child.source_parent_1_id, parent['id'])
        self.assertEqual(apps.get_model(item['model']).objects.get(pk=parent['id']).old_name, '染色')

    def test_customer_and_supplier_use_separate_scopes_on_the_same_table(self):
        customer = profile('dye.p_Customer', old_cust_kind_id='0')
        supplier = profile('dye.p_Customer', old_cust_kind_id='1')
        self.ok(self.create(customer, {'old_customer_id': '001', 'old_short_name': '同名'}))
        self.ok(self.create(supplier, {'old_num': 'SUP-1', 'old_short_name': '同名'}))
        self.assertEqual(self.count(self.client.get(self.url(customer))), 1)
        self.assertEqual(self.count(self.client.get(self.url(supplier))), 1)

    def test_unmatched_route_step_keeps_the_line_and_current_name(self):
        item = profile('dye.p_FieldValid', old_field_type='Production')
        parent = self.ok(self.create(item, {'old_code': 'PT02', 'old_name': '原路线名称'}))
        response = self.client.post('/api/system/erp_base_sourcedyepfieldvalidsub/', {'_entry': 'p_ProductionDetail', 'source_parent_1': parent['id'], 'old_sub_code': 'MISSING', 'old_sort_value': 1}, format='json')
        self.ok(response)
        self.assertEqual(apps.get_model(item['model']).objects.get(pk=parent['id']).old_name, parent['old_name'])
        line = apps.get_model('erp_base.SourceDyePFieldValidSub').objects.get()
        self.assertEqual(line.old_sub_code, 'MISSING')
        self.assertTrue(BusinessAction.objects.filter(action='entry_route_name_pending', reason__contains=str(line.pk)).exists())

    def test_unclassified_and_orphan_rows_remain_in_original_overview(self):
        item = profile('dye.p_FieldValid')
        m = apps.get_model(item['model'])
        m.objects.create(old_field_type='未登记类别', old_code='D', old_name='异常')
        m.objects.create(old_field_type='', old_code='D', old_name='异常')
        child = apps.get_model('erp_base.SourceDyePFieldValidSub')
        child.objects.create(old_field_type='Production', old_code='不存在父码', old_sub_code='GX', old_sort_value=1)
        self.assertEqual(self.count(self.client.get(self.url(item))), 2)
        load()
        self.assertEqual(m.objects.count(), 2)
        self.assertEqual(child.objects.count(), 1)

    def test_system_enum_does_not_use_same_number_lookup_as_its_source(self):
        a = profile('dye.s_DBComboValue', old_combo_id='B001')
        b = profile('dye.s_DBComboValue', old_combo_id='B003')
        self.ok(self.create(a, {'old_key_value': '0', 'old_display_value': '米'}))
        self.ok(self.create(b, {'old_key_value': '0', 'old_display_value': 'M'}))
        self.assertEqual(self.count(self.client.get(self.url(a))), 1)
        self.assertEqual(self.count(self.client.get(self.url(b))), 1)

    def test_paused_finance_profile_is_readonly(self):
        item = profile('dye.p_FieldValid', old_field_type='APKind')
        response = self.create(item, {'old_code': 'F', 'old_name': '财务暂缓'})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(apps.get_model(item['model']).objects.count(), 0)

    def test_viewer_gets_metadata_and_lists_but_no_new_write_permission(self):
        viewer = get_user_model().objects.create(username='foundation-viewer', dept=self.actor.dept)
        viewer.role.add(Role.objects.get(key='erp_original_viewer'))
        self.client.force_authenticate(viewer)
        item = profile('dye.WidthGramWeight', old_kind='幅宽')
        response = self.client.get('/api/system/' + item['entity'] + '/foundation_entry/', {'key': item['key']})
        self.ok(response)
        self.assertEqual(self.client.get(self.url(item)).status_code, 200)
        self.assertEqual(self.create(item, {'old_code': 'x'}).status_code, 403)
        menu = Menu.objects.get(component_name=item['key'])
        self.assertTrue(RoleMenuPermission.objects.filter(role__key='erp_original_viewer', menu=menu).exists())

    def history(self, item, values, scope=True):
        m = apps.get_model(item['model'])
        rows = [{key: r.get(old) for old, key in m.SOURCE_FIELDS.items()} for r in values]
        stream = BytesIO()
        write_workbook(rows, columns(m), stream, 'create', {'导入方式': '历史基础资料原值承接'})
        url = self.url(item, 'import/') if scope else '/api/system/' + item['entity'] + '/import/'
        return self.client.post(url, {'file': SimpleUploadedFile('原值.xlsx', stream.getvalue())}, format='multipart')

    def test_scoped_historical_import_preserves_all_values_and_duplicate_rows_without_events(self):
        item = profile('dye.WidthGramWeight', old_kind='幅宽')
        response = self.history(item, [{'Code': '004', 'NameC': 'CM', 'Kind': '幅宽'}, {'Code': '004', 'NameC': 'CM', 'Kind': '幅宽'}])
        self.ok(response)
        m = apps.get_model(item['model'])
        self.assertEqual(m.objects.count(), 2)
        self.assertTrue(all((r.old_num_value is None and stage(r) is None for r in m.objects.all())))
        self.assertEqual(LegacyImportBatch.objects.get().row_count, 2)
        self.assertFalse(SourceDocumentContext.objects.exists())
        self.assertFalse(BusinessAction.objects.exists())

    def test_mixed_historical_categories_roll_back_and_original_overview_can_keep_them(self):
        item = profile('dye.WidthGramWeight', old_kind='幅宽')
        data = [{'Code': '004', 'NameC': 'CM', 'Kind': '幅宽'}, {'Code': '005', 'NameC': '异常原单位', 'Kind': None}]
        self.assertEqual(self.history(item, data).status_code, 400)
        self.assertFalse(apps.get_model(item['model']).objects.exists())
        self.assertFalse(LegacyImportBatch.objects.exists())
        self.ok(self.history(item, data, scope=False))
        self.assertEqual(apps.get_model(item['model']).objects.count(), 2)

    def test_new_template_keeps_scope_when_imported_from_original_overview(self):
        item = profile('dye.WidthGramWeight', old_kind='幅宽')
        response = self.client.get(self.url(item, 'import_template/'))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(b''.join(response.streaming_content)))
        settings = dict(book['导入设置'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(settings['基础资料入口'], item['key'])
        headers = [c.value for c in book['数据'][1]]
        m = apps.get_model(item['model'])
        for old, value in [('Code', '004'), ('NameC', 'CM')]:
            title = str(m._meta.get_field(m.SOURCE_FIELDS[old]).verbose_name)
            book['数据'].cell(2, headers.index(title) + 1).value = value
        stream = BytesIO()
        book.save(stream)
        book.close()
        result = self.client.post('/api/system/' + item['entity'] + '/import/', {'file': SimpleUploadedFile('新增.xlsx', stream.getvalue())}, format='multipart')
        self.ok(result)
        self.assertEqual(m.objects.get().old_kind, '幅宽')
