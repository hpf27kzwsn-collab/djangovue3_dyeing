"""客户本地数量集用于质量统计，重复、空值及负数原样保留。"""
from plugins.erp_base.tests.source_fixtures import historical_batch
from datetime import datetime
from decimal import Decimal
from django.apps import apps
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient
from system.models import Users
from plugins.erp_quality.quantity_models import QualityQuantityBatch,QualityQuantitySource
from plugins.erp_quality.quantity_services import automatic
from plugins.erp_base.source_rework_extended import combined,responsibility
from plugins.erp_base.source_chain_reverse import reader
from plugins.erp_base.source_outwork_material import receipt_options,receive_material,return_to_factory
from plugins.erp_base.source_outwork_flow import dispatch
from plugins.erp_base.source_outwork_close import close,closure
from plugins.erp_base.source_outwork_followup import review
from plugins.erp_base.source_purchase_flow import raw_receipts,receipt_progress,receive
from plugins.erp_base.source_continuation import view_effects

class QuantityMappingTests(TestCase):
    def setUp(self):
        self.actor=Users.objects.create_superuser(username='quantity-map',password='test')
        self.client=APIClient();self.client.force_authenticate(self.actor)
        self.at=timezone.make_aware(datetime(2025,8,15,12))
    def row(self,**extra):
        return dict(detail_no='D',header_sheet_series='H',sheet_date=self.at,sheet_kind=30,
            job_number='J',quantity=Decimal('1000.00'),zero_amount=True,**extra)
    def local(self,rows,name='客户数量集'):
        batch=QualityQuantityBatch.objects.create(name=name)
        for row in rows:QualityQuantitySource.objects.create(batch=batch,**row)
        return batch
    def mapped(self,**kwargs):
        return automatic(self.at,self.at,QualityQuantityBatch.objects.all(),QualityQuantitySource.objects.all(),**kwargs)
    def test_local_duplicates_negative_null_and_unresolved_codes(self):
        base=self.row()
        batch=self.local([{**base,'quantity':-2},{**base,'quantity':None,'zero_amount':None},
            {**base,'detail_no':'ORPHAN','header_sheet_series':None,'sheet_date':None,'sheet_kind':99,'zero_amount':False}])
        rows=list(batch.quantity_rows.order_by('pk'))
        self.assertEqual(len(rows),3);self.assertEqual(rows[0].detail_no,rows[1].detail_no)
        self.assertEqual(rows[0].quantity,-2);self.assertIsNone(rows[1].quantity);self.assertIsNone(rows[1].zero_amount)
        self.assertFalse(rows[2].zero_amount);self.assertIsNone(rows[2].header_sheet_series)
    def test_same_name_groups_independent_and_explicit_selection(self):
        first=self.local([self.row()]);second=self.local([{**self.row(),'quantity':2000}])
        rows,_,status=self.mapped();self.assertEqual(rows[0]['quantity'],2000);self.assertEqual(status['batch_id'],second.pk)
        rows,_,_=self.mapped(batch=first);self.assertEqual(rows[0]['quantity'],1000)
    def test_empty_latest_group_does_not_fall_back(self):
        self.local([self.row()]);batch=self.local([])
        rows,_,status=self.mapped();self.assertEqual(rows,[]);self.assertEqual(status['batch_id'],batch.pk)
    def test_both_years_and_unknown_dates_without_duplicate_range(self):
        old=self.at.replace(year=2024);base=self.row()
        self.local([base,{**base,'sheet_date':old},{**base,'sheet_date':None},{**base,'sheet_date':self.at.replace(year=2023)}])
        rows,_,_=self.mapped(previous_start=old,previous_end=old)
        self.assertEqual(len(rows),3);self.assertEqual(len({r['source_id'] for r in rows}),3)
    def test_scope_is_respected(self):
        batch=self.local([self.row(),self.row()]);last=batch.quantity_rows.latest('pk')
        rows,_,status=automatic(self.at,self.at,QualityQuantityBatch.objects.all(),QualityQuantitySource.objects.filter(pk=last.pk))
        self.assertEqual(len(rows),1);self.assertEqual(rows[0]['source_id'],last.pk);self.assertEqual(status['batch_rows'],1)
    def test_missing_group(self):
        rows,_,status=self.mapped();self.assertEqual(rows,[]);self.assertEqual(status['state'],'unavailable')
    def test_reports_join_actual_business_models(self):
        self.local([self.row(),{**self.row(),'quantity':-500}])
        apps.get_model('erp_production.SourceDyeJobSheet').objects.create(old_sheet_series='J',old_sheet_no='J',old_loss_str='Q')
        apps.get_model('erp_quality.SourceDyeWDefect').objects.create(old_sheet_series='Q',old_sheet_no='Q',old_defect_dept='07',old_defect_class_id='03')
        rows,ref,_=self.mapped()
        self.assertEqual(combined(self.at,self.at,rows,ref)['rows'][0]['FARKg'],Decimal('.50'))
        result=responsibility(self.at,self.at,rows,ref)['rows'][0]
        self.assertEqual(result['JQty1'],Decimal('.50'));self.assertIsNone(result['JMeter1'])
    def test_actual_length_keeps_precision(self):
        batch=self.local([{**self.row(),'metres':Decimal('-123.456789')}])
        self.assertEqual(batch.quantity_rows.get().metres,Decimal('-123.456789'))
    def test_api_automatic_and_manual_sources(self):
        batch=self.local([self.row()]);url='/api/system/erp_quality_sourcedyewdefect/rework_extended/'
        args={'start':self.at.isoformat(),'end':self.at.isoformat(),'report':'Sp_ReworkCloth4'}
        response=self.client.post(url,args,format='json');self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.data['data']['rows'][0]['FARKg'],'1.00')
        self.assertEqual(response.data['data']['quantity_mapping']['batch_id'],batch.pk)
        response=self.client.post(url,{**args,'quantity_source':'manual','far_rows':[]},format='json')
        self.assertIsNone(response.data['data']['rows'][0]['FARKg'])
    def test_quantity_resources_support_customer_crud(self):
        root='/api/system/erp_quality_qualityquantitybatch/'
        response=self.client.post(root,{'name':'手工数量'},format='json');self.assertEqual(response.status_code,200,response.data)
        pk=response.data['data']['id']
        line='/api/system/erp_quality_qualityquantitysource/'
        for qty in ['-1.00',None,'0.00']:
            response=self.client.post(line,{'batch':pk,'quantity':qty,'detail_no':'SAME'},format='json')
            self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(QualityQuantitySource.objects.filter(batch_id=pk).count(),3)
        response=self.client.patch(root+str(pk)+'/',{'name':'改名'},format='json');self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(QualityQuantityBatch.objects.get(pk=pk).name,'改名')
    def test_excel_numeric_group_id_keeps_duplicate_rows_and_rejects_unknown_id_atomically(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from django.core.files.uploadedfile import SimpleUploadedFile
        batch=self.local([])
        root='/api/system/erp_quality_qualityquantitysource/'
        template=self.client.get(root+'import_template/')
        content=b''.join(template.streaming_content) if template.streaming else template.content
        book=load_workbook(BytesIO(content));headers={cell.value:cell.column for cell in book['数据'][1]}
        for number in [2,3]:
            for title,value in {'数量集':batch.pk,'来源明细号':'SAME','数量':-2.5}.items():book['数据'].cell(number,headers[title],value)
        def upload():
            out=BytesIO();book.save(out)
            return self.client.post(root+'import/',{'file':SimpleUploadedFile('quantity.xlsx',out.getvalue())},format='multipart')
        response=upload();self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(list(batch.quantity_rows.values_list('quantity',flat=True)),[Decimal('-2.50')]*2)
        book['数据'].cell(3,headers['数量集'],999999)
        response=upload();self.assertEqual(response.status_code,400,response.data)
        self.assertEqual(batch.quantity_rows.count(),2)
        book.close()


class MaterialContinuationTests(TestCase):
    def setUp(self):
        self.actor=Users.objects.create_superuser(username='material-cont',password='test')
        self.at=timezone.make_aware(datetime(2026,10,15,12));self.client=APIClient();self.client.force_authenticate(self.actor)
    def old(self,name,module='erp_warehouse',**values):
        row=apps.get_model(module+'.SourceDye'+name).objects.create(**values);row.refresh_from_db();return row
    def source(self,family='Gray',**values):
        head=self.old(family+'Out',old_sheet_series='OUT'+family,old_sheet_kind=30 if family=='Gray' else 32,old_stock_unit=1)
        return self.old(family+'OutDetail',old_sheet_series=head.old_sheet_series,old_assigned_no='O'+family,
            source_parent_1=head,**{'old_weight_qty':10,'old_rolls':2,'old_packs':3,'old_job_sheet_no':'J',**values})
    def back(self,table,kind,**values):
        head=self.old(table,old_sheet_series=table+str(kind),old_sheet_kind=kind,old_stock_unit=1)
        return self.old(table+'Detail',old_sheet_series=head.old_sheet_series,old_rece_sheet_no=table+str(kind),source_parent_1=head,**values)
    def actual(self,item,destination='semi',**extra):
        return receive_material(item,self.actor,destination,weight=-2,length=None,rolls=Decimal('.5'),packs=None,
            warehouse='W',unit=1,customer='C',supplier='S',at=self.at,reason='实际加工收回',receipt_number='NEW-'+destination,**extra)

    def test_gray_original_three_balances_and_mixed_roll_formula(self):
        row=self.source()
        self.back('SemiIn',25,old_job_sheet_no='J',old_rolls=2,old_weight_qty=3,old_packs=1)
        self.back('GrayIn',21,old_cust_job_no='J',old_rolls=1,old_weight_qty=4,old_packs=1)
        result=receipt_options(row)['rows']
        self.assertEqual([line['remaining_weight'] for line in result],[3,3])
        self.assertEqual([line['original_rolls_formula'] for line in result],[4,7])
        self.assertTrue(all(line['old_selection_available'] for line in result))

    def test_semi_26_reference_and_original_weight_named_rolls(self):
        row=self.source('Semi',old_detail_no='D')
        self.back('SemiIn',26,old_gray_out_detail_no='D',old_weight_qty=12,old_rolls=1)
        result=receipt_options(row)['rows'][0]
        self.assertEqual(result['original_rolls_formula'],-2);self.assertEqual(result['remaining_weight'],-2)
        self.assertFalse(result['old_selection_available'])

    def test_original_null_vs_blank_reference_semantics(self):
        row=self.source(old_job_sheet_no=None)
        self.back('SemiIn',25,old_job_sheet_no=None,old_weight_qty=4)
        self.assertEqual(receipt_options(row)['rows'][0]['remaining_weight'],10)
        row.old_job_sheet_no='';row.save()
        self.assertEqual(receipt_options(row)['rows'][0]['remaining_weight'],6)

    def test_duplicate_heads_expand_and_other_kind_is_visible(self):
        row=self.source()
        line=self.back('SemiIn',25,old_job_sheet_no='J',old_weight_qty=3)
        self.old('SemiIn',old_sheet_series=line.old_sheet_series,old_sheet_kind=25)
        self.old('GrayOut',old_sheet_series=row.old_sheet_series,old_sheet_kind=99)
        result=receipt_options(row)
        self.assertEqual(len(result['result_rows']),2);self.assertEqual(len(result['rows']),4)
        self.assertTrue(any(not r['old_selection_available'] for r in result['rows']))

    def test_receive_all_three_variants_actual_fields_and_stock(self):
        for family,destination,kind in [('Gray','gray',21),('Gray','semi',25),('Semi','semi',26)]:
            source=self.source(family,**({'old_detail_no':'D'} if family=='Semi' else {}))
            target=self.actual(source,destination)
            self.assertEqual(target.old_sheet_kind,kind)
            line=apps.get_model('erp_warehouse.SourceDye'+target.SOURCE_TABLE+'Detail').objects.get(source_parent_1=target)
            self.assertEqual(line.old_weight_qty,-2);self.assertEqual(line.old_rolls,Decimal('.5'));self.assertIsNone(line.old_length_qty)
            if kind==21:self.assertEqual(line.old_cust_job_no,'J')
            elif kind==25:self.assertEqual(line.old_job_sheet_no,'J')
            else:self.assertEqual(line.old_gray_out_detail_no,'D')
            book=apps.get_model('erp_warehouse.SourceDye'+target.SOURCE_TABLE.removesuffix('In')+'StockBook').objects.get(old_sheet_series=target.old_sheet_series)
            self.assertEqual(book.old_in_weight_qty,-2)
            self.assertEqual(view_effects(target)['outwork_receipt_origin']['source']['id'],source.pk)

    def test_after_closure_material_receipt_marks_review_and_reclose_includes_both_models(self):
        app=self.old('OutWorkApply','erp_production',old_sheet_series='A')
        apply=self.old('OutWorkApplyDetail','erp_production',old_sheet_series='A',old_detail_no='A1',source_parent_1=app)
        outhead=dispatch(apply,self.actor,10,1,'W',1,30,self.at,'J')
        source=apps.get_model('erp_warehouse.SourceDyeGrayOutDetail').objects.get(source_parent_1=outhead)
        close(app,self.actor,[source],[],[],self.at,'初次')
        received=self.actual(source)
        self.assertEqual(closure(app)['state'],'needs_review')
        semi=apps.get_model('erp_warehouse.SourceDyeSemiInDetail').objects.get(source_parent_1=received)
        grayhead=self.old('GrayIn',old_sheet_series='GR',old_stock_unit=1)
        gray=self.old('GrayInDetail',id=semi.pk,old_rece_sheet_no='GR',old_sheet_series='GR',source_parent_1=grayhead,old_weight_qty=4)
        result=close(app,self.actor,[source],[],[],self.at,'复核',received_gray=[gray],received_semi=[semi])
        weight=next(row for row in result['rows'] if row['measure']=='WeightQty')
        self.assertEqual(weight['received'],2);self.assertEqual(weight['difference'],8)
        self.assertEqual(len(result['sources']),3);self.assertEqual(review(app,reader)['closure_versions'],2)

    def test_purchase_reports_both_references_without_double_count_or_rewriting_progress(self):
        item=self.old('PurDetail','erp_purchase',old_detail_no='P')
        self.back('GrayIn',18,old_purchase_orders='P',old_pur_detail_no='OTHER',old_weight_qty=3)
        self.back('SemiIn',20,old_purchase_orders='P',old_pur_detail_no='P',old_weight_qty=4)
        report=raw_receipts(item,reader)
        self.assertEqual(len(report['rows']),2)
        self.assertEqual(report['rows'][0]['purchase_reference_fields'],['PurchaseOrders'])
        self.assertEqual(report['rows'][1]['purchase_reference_fields'],['PurDetailNo','PurchaseOrders'])
        self.assertEqual(receipt_progress(item,reader)['received_weight_18_20'],4)

    def test_purchase_actual_length_supplier_non_numeric_number_both_families(self):
        item=self.old('PurDetail','erp_purchase',old_detail_no='P',old_price=2)
        for entry,table in [('PurGrayIn','GrayIn'),('PurSemiIn','SemiIn')]:
            self.old(table+'Detail',old_rece_sheet_no='NONNUMERIC',old_sheet_series='OLD')
            head=receive(item,self.actor,entry,self.at,Decimal('-2'),None,'W',1,length=Decimal('3.25'),supplier='S',receipt_number='NEW')
            line=apps.get_model('erp_warehouse.SourceDye'+table+'Detail').objects.get(source_parent_1=head)
            self.assertEqual(line.old_length_qty,Decimal('3.25'));self.assertEqual(head.old_supplier_id,'S')
            self.assertEqual(line.old_pur_detail_no,'P');self.assertEqual(line.old_purchase_orders,'P')

    def test_material_receipt_api_and_balance(self):
        row=self.source();root=f'/api/system/erp_warehouse_sourcedyegrayoutdetail/{row.pk}/'
        self.assertEqual(self.client.post(root+'material_receipt_options/',{},format='json').status_code,200)
        response=self.client.post(root+'receive_outwork_material/',{'destination':'semi','weight':'-2','at':self.at.isoformat(),
            'reason':'真实回料','receipt_number':'NEW'},format='json')
        self.assertEqual(response.status_code,200,response.data)

    def test_factory_disposition_for_both_material_families_and_actual_return_preserves_receipt(self):
        for family in ['Gray','Semi']:
            source=self.back(family+'In',21 if family=='Gray' else 26,old_weight_qty=5,old_job_sheet_no='J')
            url=f'/api/system/erp_warehouse_sourcedye{family.lower()}indetail/{source.pk}/factory_quality/'
            response=self.client.post(url,{'result':'hold','factory':'F','at':self.at.isoformat(),'reason':'复检'},format='json')
            self.assertEqual(response.status_code,200,response.data)
            target=return_to_factory(source,self.actor,-1,None,None,None,'W',1,'C','F',32,self.at,'实际退厂',issue_number='RETURN')
            source.refresh_from_db();self.assertEqual(source.old_weight_qty,5)
            self.assertEqual(view_effects(source)['factory_quality']['result'],'hold')
            self.assertEqual(view_effects(target)['outwork_return_origin']['factory'],'F')
            line=apps.get_model('erp_warehouse.SourceDye'+family+'OutDetail').objects.get(source_parent_1=target)
            self.assertEqual(line.old_rece_sheet_no,source.old_rece_sheet_no);self.assertEqual(line.old_weight_qty,-1)

    def test_new_material_receipt_and_factory_return_follow_execution_cancel_and_undo(self):
        from plugins.erp_base.source_chain_reverse import discover,cancel,undo
        from plugins.erp_base.source_execution_state import active
        job=self.old('JobSheet','erp_production',old_sheet_series='J',old_sheet_no='J')
        source=self.source()
        receipt=self.actual(source)
        line=apps.get_model('erp_warehouse.SourceDyeSemiInDetail').objects.get(source_parent_1=receipt)
        line.old_job_sheet_no=None;line.save()  # 没有工卡号仍通过明确办理来源追溯。
        target=return_to_factory(line,self.actor,1,None,None,None,'W',1,'C','F',32,self.at,'实际退厂',issue_number='RETURN')
        nodes=discover(job)['nodes'];self.assertIn((receipt._meta.label,receipt.pk),nodes);self.assertIn((target._meta.label,target.pk),nodes)
        cancel(job,self.actor,'execution_only','trigger_then_maintenance',self.at,'执行撤销')
        self.assertFalse(active(type(target).objects.filter(pk=target.pk)).exists())
        undo(job,self.actor,self.at,'恢复')
        self.assertTrue(active(type(target).objects.filter(pk=target.pk)).exists())

    def test_failed_material_return_rolls_back_new_header_and_lines(self):
        source=self.back('GrayIn',21,old_weight_qty=3)
        head=apps.get_model('erp_warehouse.SourceDyeGrayOut')
        before=head.objects.count()
        with self.assertRaises(ValidationError):
            return_to_factory(source,self.actor,Decimal('1000000000000'),None,None,None,'W',1,'C','F',30,self.at,'超字段精度',issue_number='RETURN')
        self.assertEqual(head.objects.count(),before)

    def test_quality_disposition_after_closure_marks_followup_without_stock(self):
        from plugins.erp_base.source_delivery_flow import factory_quality
        app=self.old('OutWorkApply','erp_production',old_sheet_series='A')
        line=self.back('SemiIn',25,old_weight_qty=3)
        close(app,self.actor,[],[],[],self.at,'结案',received_semi=[line])
        book=apps.get_model('erp_warehouse.SourceDyeSemiStockBook')
        before=book.objects.count()
        factory_quality(line,self.actor,'F','return',None,self.at,'原质量退回')
        self.assertEqual(closure(app)['state'],'needs_review');self.assertEqual(book.objects.count(),before)

    def test_typed_purchase_tables_all_receipt_entries_and_source_units(self):
        from plugins.erp_base.source_purchase_flow import receive_material as purchase_receive,material_receipts
        for family,entries in [('Stuff',['StuffIn','SmallStuffIn']),('Article',['ArticleIn','FabricArticleIn'])]:
            parent=self.old(family+'PurMain','erp_purchase',old_sheet_series=family)
            item=self.old(family+'PurDetail','erp_purchase',old_sheet_series=family,old_detail_no=family,
                old_qty=10,old_need_qty=12,old_pack_unit='BOX',old_price=2,
                **({'old_price_unit':3,'old_stuff_id':'S'} if family=='Stuff' else {'old_unit':'PCS','old_article_id':'A'}))
            for entry in entries:
                head=purchase_receive(item,self.actor,entry,Decimal('-2'),'W','SUP',15,self.at,'实际采购收货')
                line=apps.get_model('erp_warehouse.SourceDye'+family+'InDetail').objects.get(source_parent_1=head)
                self.assertEqual(line.old_pur_detail_no,family)
                self.assertEqual(getattr(line,'old_qty' if family=='Stuff' else 'old_in_qty'),-2)
                origin=view_effects(head)['purchase_origin']
                self.assertEqual(origin['source']['id'],item.pk)
                self.assertEqual(origin['headers'][0]['entity'],parent._meta.label)
            item.refresh_from_db();self.assertEqual(item.old_qty,10);self.assertEqual(item.old_need_qty,12)
            report=material_receipts(item,reader);self.assertEqual(len(report['rows']),2)
            self.assertTrue(all(row['source_unit']==(None if family=='Stuff' else 'PCS') for row in report['rows']))
            self.assertTrue(all(row['source_price_unit']==(3 if family=='Stuff' else None) for row in report['rows']))

    def test_shared_purchase_original_headers_retained_and_api_available(self):
        from plugins.erp_base.source_purchase_flow import receive_material as purchase_receive
        for parent in ['PurMain','StuffPurMain','ArticlePurMain']:
            self.old(parent,'erp_purchase',old_sheet_series='SHARED')
        item=self.old('PurDetail','erp_purchase',old_sheet_series='SHARED',old_detail_no='P',old_qty=10,old_stuff_unit='KG')
        head=purchase_receive(item,self.actor,'StuffIn',Decimal('2'),'W','S',15,self.at,'实际收料')
        self.assertEqual(len(view_effects(head)['purchase_origin']['headers']),3)
        for family in ['PurDetail','StuffPurDetail','ArticlePurDetail']:
            source=item if family=='PurDetail' else self.old(family,'erp_purchase',old_detail_no=family)
            response=self.client.post(f'/api/system/erp_purchase_sourcedye{family.lower()}/{source.pk}/material_receipts/',{},format='json')
            self.assertEqual(response.status_code,200,response.data)

    def test_typed_purchase_wrong_material_family_does_not_create_receipt(self):
        from plugins.erp_base.source_purchase_flow import receive_material as purchase_receive
        source=self.old('StuffPurDetail','erp_purchase',old_detail_no='S')
        with self.assertRaises(ValidationError):
            purchase_receive(source,self.actor,'ArticleIn',Decimal('2'),'W','S',15,self.at,'原物料家族不符')
        self.assertEqual(apps.get_model('erp_warehouse.SourceDyeArticleIn').objects.count(),0)
