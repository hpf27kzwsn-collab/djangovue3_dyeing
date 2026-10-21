"""通过实际xlsx接口验证入口、编号、父子保存、空值及整批回滚。"""
import json
from io import BytesIO
from zipfile import ZipFile
from django.test import TestCase
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient
from plugins.erp_base.source_phase3_operations import model
from plugins.erp_base.source_entry_services import save_context
from plugins.erp_base.source_context import SourceDocumentContext
from plugins.erp_base.source_sequences import SourceIdentitySequence
from plugins.erp_base.models import BusinessAction
from plugins.erp_base import source_services
from system.excel_fields import UPDATE_KEY,SERIAL
from .test_phase_report_audit import create


class SourceExcelTests(TestCase):
    def setUp(self):
        self.actor=get_user_model().objects.create(username='excel-entry',is_superuser=True)
        self.client=APIClient();self.client.force_authenticate(self.actor)

    def url(self,table,action):return '/api/system/'+model(table)._meta.label_lower.replace('.','_')+'/'+action+'/'

    def template(self,table,entry=None,update=False):
        response=self.client.get(self.url(table,'update_template' if update else 'import_template'),{'_entry':entry} if entry else {})
        self.assertEqual(response.status_code,200,getattr(response,'data',None))
        return load_workbook(BytesIO(b''.join(response.streaming_content)))

    def fill(self,book,table,records):
        sheet=book['数据'];headers=[c.value for c in sheet[1]]
        titles={}
        for name,field in model(table).SOURCE_FIELDS.items():
            caption=str(model(table)._meta.get_field(field).verbose_name)
            titles[name]=caption+' ['+name+']' if caption+' ['+name+']' in headers else caption
        for number,record in enumerate(records,2):
            values={titles.get(k,k):v for k,v in record.items()}
            for col,title in enumerate(headers,1):sheet.cell(number,col).value=values.get(title)
            self.assertFalse(set(values)-set(headers),set(values)-set(headers))
        return book

    def upload(self,table,book,entry=None):
        stream=BytesIO();book.save(stream);book.close()
        data={'file':SimpleUploadedFile('录入.xlsx',stream.getvalue())}
        if entry:data['_entry']=entry
        return self.client.post(self.url(table,'import'),data,format='multipart')

    def ok(self,response):self.assertEqual(response.status_code,200,response.data);return response.data['data']

    def test_customer_duplicate_original_codes_run_formula_and_preserve_text(self):
        book=self.fill(self.template('p_Customer'),'p_Customer',[
            {'业务入口':'p_Customer','CustomerID':'0012','ShortName':'  重复客户  ','CustKindID':'0'},
            {'业务入口':'p_Customer','CustomerID':'0012','ShortName':'#TEXT:#NULL','Num':'#NULL'}])
        result=self.ok(self.upload('p_Customer',book))
        rows=list(model('p_Customer').objects.order_by('pk'))
        self.assertEqual([(r.old_customer_id,r.old_short_name,r.old_num,r.entry_name) for r in rows],
                         [('0012','  重复客户  ','K0012','p_Customer'),('0012','#NULL',None,'p_Customer')])
        self.assertEqual([r['id'] for r in result['rows']],[r.pk for r in rows])
        self.assertEqual(result['skipped'],0)

    def test_template_entry_default_and_row_override_share_table_but_keep_own_rules(self):
        book=self.fill(self.template('p_Customer','p_Customer'),'p_Customer',[
            {'CustomerID':'C1'},{'业务入口':'p_Supplier','Num':'SUP-1'}])
        self.ok(self.upload('p_Customer',book,entry='p_Supplier'))
        rows=list(model('p_Customer').objects.order_by('pk'))
        self.assertEqual([(r.entry_name,r.old_customer_id,r.old_num) for r in rows],[('p_Customer','C1','KC1'),('p_Supplier','SUP-1','SUP-1')])

    def test_request_entry_and_automatic_numbers_are_persisted_per_row(self):
        book=self.fill(self.template('CookIn'),'CookIn',[{SERIAL:1,'Remarks':'同内容'},{SERIAL:2,'Remarks':'同内容'}])
        self.ok(self.upload('CookIn',book,'CookIn'))
        rows=list(model('CookIn').objects.order_by('pk'))
        self.assertEqual(len(rows),2)
        self.assertEqual(len({r.old_sheet_series for r in rows}),2)
        self.assertEqual(len({r.old_sheet_no for r in rows}),2)
        self.assertTrue(all(r.old_sheet_series and r.old_sheet_no and r.old_sheet_date and r.entry_name=='CookIn' for r in rows))

    def test_missing_entry_rolls_back_prior_numbers_context_and_rows(self):
        book=self.fill(self.template('CookIn'),'CookIn',[{'业务入口':'CookIn','Remarks':'先执行'},{'Remarks':'缺入口'}])
        response=self.upload('CookIn',book)
        self.assertEqual(response.status_code,400,response.data)
        self.assertEqual(str(response.data['data']['row']),'3')
        self.assertFalse(model('CookIn').objects.exists())
        self.assertFalse(SourceIdentitySequence.objects.exists())
        self.assertFalse(SourceDocumentContext.objects.exists())

    def test_explicit_null_empty_and_number_values_do_not_request_defaults(self):
        book=self.fill(self.template('CookIn','CookIn'),'CookIn',[
            {'SheetSeries':'DUP','SheetNo':'#NULL','SheetDate':'#NULL','Remarks':'#EMPTY','PrintTimes':0},
            {'SheetSeries':'DUP','SheetNo':'#EMPTY','SheetDate':'#NULL','Remarks':'   ','PrintTimes':-2}])
        self.ok(self.upload('CookIn',book))
        rows=list(model('CookIn').objects.order_by('pk'))
        self.assertEqual([(r.old_sheet_no,r.old_sheet_date,r.old_remarks,r.old_print_times) for r in rows],[(None,None,'',0),('',None,'   ',-2)])

    def test_update_blank_keep_null_empty_and_duplicate_targets_execute_in_order(self):
        item=source_services.write(model('p_Customer'),{'_entry':'p_Customer','old_customer_id':'C','old_short_name':'原文','old_name_c':'名称','old_remarks':'旧备注'},actor=self.actor)
        book=self.fill(self.template('p_Customer','p_Supplier',update=True),'p_Customer',[
            {UPDATE_KEY:item.pk,'ShortName':None,'NameC':'#NULL','Remarks':'#EMPTY'},
            {UPDATE_KEY:item.pk,'ShortName':'#KEEP','Remarks':'#TEXT:#EMPTY'}])
        result=self.ok(self.upload('p_Customer',book,'p_Supplier'));item.refresh_from_db()
        self.assertEqual((item.entry_name,item.old_short_name,item.old_name_c,item.old_remarks),('p_Customer','原文',None,'#EMPTY'))
        self.assertEqual(result['updated'],2)
        self.assertEqual([r['id'] for r in result['rows']],[item.pk,item.pk])

    def test_update_template_roundtrips_null_empty_and_reserved_literal_text(self):
        item=source_services.write(model('p_Customer'),{'_entry':'p_Customer','old_customer_id':'C','old_short_name':'#KEEP','old_name_c':'','old_remarks':None},actor=self.actor)
        book=self.template('p_Customer',update=True)
        values=dict(zip([c.value for c in book['数据'][1]],[c.value for c in book['数据'][2]]))
        self.assertIn('#TEXT:#KEEP',values.values());self.assertIn('#NULL',values.values());self.assertIn('#EMPTY',values.values())
        self.ok(self.upload('p_Customer',book));item.refresh_from_db()
        self.assertEqual((item.old_short_name,item.old_name_c,item.old_remarks),('#KEEP','',None))

    def test_cook_save_event_uses_all_original_receipt_candidates(self):
        for _ in range(2):create('ArticleIn',SheetSeries='COOK-X',SheetKind=13)
        lines=[create('ArticleInDetail',SheetSeries='COOK-X',ArticleID='MAT',TaxPrice=v) for v in [2,-3,None]]
        mats=[create('p_Article',ArticleID='MAT',Price=v) for v in [None,0,9]]
        book=self.fill(self.template('CookIn','CookIn'),'CookIn',[{'SheetSeries':'COOK-X','SheetNo':'X'}])
        self.ok(self.upload('CookIn',book))
        for mat,price in zip(mats,[None,None,9]):
            mat.refresh_from_db();self.assertEqual(mat.old_price,price);self.assertIsNone(mat.old_last_price)
            event=BusinessAction.objects.get(entity=mat._meta.label,object_id=mat.pk,action='entry_cook_price_candidates')
            self.assertEqual([c['detail_id'] for c in json.loads(event.reason)['candidates']],[r.pk for r in lines]*2)

    def test_parent_detail_import_inherits_entry_and_updates_head_and_stock(self):
        head=source_services.write(model('ProdIn'),{'_entry':'TwistingIn','old_sheet_series':'TIN','old_sheet_no':'TIN'},actor=self.actor)
        book=self.template('ProdInRecord')
        field=model('ProdInRecord')._meta.get_field('source_parent_1')
        self.fill(book,'ProdInRecord',[
            {str(field.verbose_name):f'{head.pk} [{head.pk}]','WeightQty':-5,'InRolls':-2},
            {str(field.verbose_name):f'{head.pk} [{head.pk}]','WeightQty':2,'InRolls':1}])
        self.ok(self.upload('ProdInRecord',book))
        rows=list(model('ProdInRecord').objects.order_by('pk'))
        self.assertTrue(all(r.entry_name=='TwistingInRecord' and r.old_sheet_series=='TIN' for r in rows))
        head.refresh_from_db();self.assertEqual((head.old_weight_qty,head.old_rolls),(-3,-1))
        stock=model('ProdStockBook').objects.get(old_sheet_series='TIN')
        self.assertEqual((stock.old_in_weight_qty,stock.old_in_rolls),(-3,-1))

    def test_nonempty_serial_only_row_is_processed_and_blank_row_is_reported(self):
        book=self.fill(self.template('CookIn','CookIn'),'CookIn',[{SERIAL:1},{},{SERIAL:3}])
        result=self.ok(self.upload('CookIn',book))
        self.assertEqual((result['created'],result['skipped']),(2,1))
        self.assertEqual([r['status'] for r in result['rows']],['created','blank','created'])

    def test_uncached_formula_reports_cell_and_rolls_back_previous_rows(self):
        book=self.fill(self.template('CookIn','CookIn'),'CookIn',[{'Remarks':'有效行'},{'Remarks':'=1+1'}])
        response=self.upload('CookIn',book)
        self.assertEqual(response.status_code,400,response.data)
        self.assertEqual(str(response.data['data']['row']),'3')
        self.assertIn('公式',str(response.data['data']))
        self.assertFalse(model('CookIn').objects.exists())

    def test_service_error_reports_row_and_rolls_back_prior_save_effects(self):
        mat=create('p_Article',ArticleID='M',Price=0)
        create('ArticleIn',SheetSeries='C',SheetKind=13)
        create('ArticleInDetail',SheetSeries='C',ArticleID='M',TaxPrice=-7)
        book=self.fill(self.template('CookIn','CookIn'),'CookIn',[{'SheetSeries':'C'},{'业务入口':'p_Customer','Remarks':'错误入口'}])
        response=self.upload('CookIn',book)
        self.assertEqual(response.status_code,400,response.data)
        mat.refresh_from_db();self.assertEqual(mat.old_price,0)
        self.assertFalse(model('CookIn').objects.exists())
        self.assertFalse(BusinessAction.objects.exists())

    def test_duplicate_original_captions_are_disambiguated_in_template(self):
        book=self.template('ProdIn')
        headers=[c.value for c in book['数据'][1]]
        self.assertEqual(len(headers),len(set(headers)))
        self.assertTrue(any('[' in h for h in headers))

    def test_unknown_column_content_is_not_silently_dropped(self):
        book=self.fill(self.template('CookIn','CookIn'),'CookIn',[{'Remarks':'主行'}])
        sheet=book['数据'];sheet.cell(2,sheet.max_column+1,'未命名异常数据')
        response=self.upload('CookIn',book)
        self.assertEqual(response.status_code,400,response.data)
        self.assertFalse(model('CookIn').objects.exists())

    def test_update_downloaded_values_do_not_override_changed_source_formula(self):
        item=source_services.write(model('p_Customer'),{'_entry':'p_Customer','old_customer_id':'001'},actor=self.actor)
        book=self.template('p_Customer',update=True)
        title=str(type(item)._meta.get_field('old_customer_id').verbose_name)
        column=next(c.column for c in book['数据'][1] if c.value==title)
        book['数据'].cell(2,column,'002')
        self.ok(self.upload('p_Customer',book));item.refresh_from_db()
        self.assertEqual((item.old_customer_id,item.old_num),('002','K002'))

    def test_cached_formula_zero_is_imported_and_not_treated_as_blank(self):
        book=self.fill(self.template('CookIn','CookIn'),'CookIn',[{'Remarks':'公式缓存','PrintTimes':'=1-1'}])
        buffer=BytesIO();book.save(buffer);book.close();output=BytesIO()
        with ZipFile(buffer) as original,ZipFile(output,'w') as result:
            for info in original.infolist():
                data=original.read(info.filename)
                if info.filename=='xl/worksheets/sheet1.xml':
                    from xml.etree import ElementTree as ET
                    root=ET.fromstring(data);ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                    cell=next(c for c in root.iter(ns+'c') if c.findtext(ns+'f')=='1-1')
                    cell.find(ns+'v').text='0'
                    data=ET.tostring(root)
                result.writestr(info,data)
        response=self.client.post(self.url('CookIn','import'),{'file':SimpleUploadedFile('公式.xlsx',output.getvalue())},format='multipart')
        self.ok(response)
        self.assertEqual(model('CookIn').objects.get().old_print_times,0)

    def test_business_service_error_has_row_field_and_rolls_back_updates(self):
        heads=[]
        for series in ['FIRST','SECOND']:
            head=create('CookIn',SheetSeries=series,SheetNo=series,Remarks='原备注')
            head.approval_stage='draft';head.save();save_context(head,self.actor,'CookIn');heads.append(head)
        book=self.fill(self.template('CookIn',update=True),'CookIn',[
            {UPDATE_KEY:heads[0].pk,'Remarks':'不应保存'},
            {UPDATE_KEY:heads[1].pk,'SheetSeries':'CHANGE'}])
        response=self.upload('CookIn',book)
        self.assertEqual(response.status_code,400,response.data)
        self.assertEqual(str(response.data['data']['row']),'3')
        self.assertIn('表单流水号',str(response.data['data']['field']))
        for head in heads:
            head.refresh_from_db();self.assertEqual(head.old_remarks,'原备注')
        self.assertEqual(heads[1].old_sheet_series,'SECOND')
