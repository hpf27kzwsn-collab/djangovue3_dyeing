from decimal import Decimal
from io import BytesIO
from django.core.management import call_command
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient
from system.models import Users, Role, Menu, MenuButton, RoleMenuButtonPermission,RoleMenuPermission, Dept, FieldPermission
from system.catalog import ENTITIES
from .. import models as m
from ..choices import VersionStatus
from ..resource_options import MODEL_FIELDS


class MasterDataApiTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('init_data', verbosity=0)
        cls.root = Users.objects.get(username='superadmin')
        cls.kg = m.Unit.objects.create(code='kg', name='千克', dimension='mass', factor_to_base=1, decimal_places=3)
        cls.factory = m.Factory.objects.create(code='F1', name='第一工厂')
        cls.party = m.Party.objects.create(code='P1', name='客户一')
        cls.material = m.Material.objects.create(code='M1', name='纱线一', kind='yarn', stock_unit=cls.kg)
        cls.process = m.Process.objects.create(code='GX03', name='装缸')
        cls.route = m.Route.objects.create(code='R1', name='染色路线')

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.root)

    def url(self, model, pk=None, action=None):
        return '/api/system/erp_base_' + model.lower() + '/' + (str(pk) + '/' if pk is not None else '') + (action + '/' if action else '')

    def create(self, model, data, expected=200):
        response = self.client.post(self.url(model), data, format='json')
        self.assertEqual(response.status_code, expected, response.data)
        return response.data['data']

    def patch(self, model, pk, data, expected=200):
        response = self.client.patch(self.url(model, pk), data, format='json')
        self.assertEqual(response.status_code, expected, response.data)
        return response

    def version(self):
        return m.RouteVersion.objects.create(route=self.route, version=1)

    def step(self, version, sequence=1, predecessor=None):
        return m.RouteStep.objects.create(route_version=version, sequence=sequence, process=self.process,
                                         is_required=True, predecessor=predecessor)

    def test_all_resources_and_metadata_use_actual_fields(self):
        for name, fields in MODEL_FIELDS.items():
            with self.subTest(model=name):
                model = ENTITIES['erp_base_' + name.lower()]
                expected = {'id', *fields.split(), 'remarks', 'legacy_id', 'create_datetime', 'update_datetime'}
                response = self.client.get('/api/meta/erp_base_' + name.lower() + '/')
                self.assertEqual(response.status_code, 200, response.data)
                actual = {field['key'] for field in response.data['data']['fields']}
                self.assertEqual(expected, actual)
                self.assertTrue(expected <= {field.name for field in model._meta.fields})
                self.assertEqual(self.client.get(self.url(name)).status_code, 200)
        self.assertEqual(self.client.get('/api/system/erp_sales_salesorder/').status_code, 200)

    def test_legacy_identity_and_route_status_are_read_only(self):
        self.create('Party', {'code': 'P2', 'name': '客户二', 'legacy_id': '0002'}, 400)
        version = self.version()
        self.patch('RouteVersion', version.pk, {'status': 'active'}, 400)
        self.patch('Party', self.party.pk, {'legacy_id': ''}, 400)
        self.patch('ColorStandard', self.create('ColorStandard', {'code': 'C1', 'color_number': '001'})['id'], {'source_sample': None}, 400)
        row = self.create('Party', {'code': 'P2', 'name': '客户二'})
        self.assertIsNone(row['legacy_id'])

    def test_preserves_null_false_zero_empty_and_text_spaces(self):
        row = self.create('Party', {'code': 'P2', 'name': ' 客户二 ', 'phone': '', 'fax': None,
                                   'is_active': None, 'orders_enabled': False})
        self.assertEqual(row['name'], ' 客户二 ')
        self.assertEqual(row['phone'], '')
        self.assertIsNone(row['fax'])
        self.assertIsNone(row['is_active'])
        self.assertIs(row['orders_enabled'], False)
        self.patch('Party', row['id'], {'phone': None})
        row = self.create('PackagingSpecification', {'code': 'B1', 'name': '袋', 'kind': 'bag', 'unit_tare_kg': '0'})
        self.assertEqual(Decimal(row['unit_tare_kg']), 0)

    def test_default_contact_switch_and_failed_save_rollback(self):
        first = self.create('PartyContact', {'party': self.party.pk, 'name': '甲', 'is_default': True})
        second = self.create('PartyContact', {'party': self.party.pk, 'name': '乙', 'is_default': True})
        self.assertFalse(m.PartyContact.objects.get(pk=first['id']).is_default)
        self.patch('PartyContact', first['id'], {'is_default': True, 'name': 'x' * 51}, 400)
        self.assertTrue(m.PartyContact.objects.get(pk=second['id']).is_default)
        self.patch('PartyContact', first['id'], {'is_default': True})
        self.assertFalse(m.PartyContact.objects.get(pk=second['id']).is_default)

    def test_default_addresses_are_separate_per_type(self):
        a = self.create('PartyAddress', {'party': self.party.pk, 'kind': 'delivery', 'address': '甲', 'is_default': True})
        b = self.create('PartyAddress', {'party': self.party.pk, 'kind': 'registered', 'address': '乙', 'is_default': True})
        self.create('PartyAddress', {'party': self.party.pk, 'kind': 'delivery', 'address': '丙', 'is_default': True})
        self.assertFalse(m.PartyAddress.objects.get(pk=a['id']).is_default)
        self.assertTrue(m.PartyAddress.objects.get(pk=b['id']).is_default)
        other = m.Party.objects.create(code='P2', name='客户二')
        contact = m.PartyContact.objects.create(party=other, name='其他客户联系人')
        self.patch('PartyAddress', a['id'], {'is_default': True, 'contact': contact.pk}, 400)
        self.assertTrue(m.PartyAddress.objects.get(party=self.party, kind='delivery', is_default=True).is_default)

    def test_contact_and_detail_parent_membership(self):
        other = m.Party.objects.create(code='P2', name='另一个客户')
        contact = m.PartyContact.objects.create(party=other, name='外部联系人')
        self.create('PartyAddress', {'party': self.party.pk, 'kind': 'delivery', 'address': '地址', 'contact': contact.pk}, 400)
        self.patch('PartyContact', contact.pk, {'party': self.party.pk}, 400)

    def test_party_and_dictionary_parent_cycles(self):
        child = m.Party.objects.create(code='P2', name='下级', parent=self.party)
        self.patch('Party', self.party.pk, {'parent': child.pk}, 400)
        a = m.BusinessChoice.objects.create(namespace='A', code='1', name='甲')
        b = m.BusinessChoice.objects.create(namespace='A', code='2', name='乙', parent=a)
        self.patch('BusinessChoice', a.pk, {'parent': b.pk}, 400)
        self.patch('BusinessChoice', a.pk, {'namespace': 'B'}, 400)

    def test_material_specs_match_kind_and_units_keep_meaning(self):
        self.create('ChemicalSpecification', {'material': self.material.pk}, 400)
        self.create('YarnSpecification', {'material': self.material.pk, 'structure': None})
        self.patch('Material', self.material.pk, {'kind': 'chemical'}, 400)
        self.patch('Unit', self.kg.pk, {'dimension': 'length'}, 400)
        metre = m.Unit.objects.create(code='m', name='米', dimension='length', factor_to_base=1, decimal_places=2)
        self.create('Unit', {'code': 'bad', 'name': '错误', 'dimension': 'mass', 'base_unit': metre.pk, 'factor_to_base': '0.001', 'decimal_places': 3}, 400)
        gram = self.create('Unit', {'code': 'g', 'name': '克', 'dimension': 'mass', 'base_unit': self.kg.pk, 'factor_to_base': '0.001', 'decimal_places': 3})
        self.assertEqual(m.Unit.objects.get(pk=gram['id']).convert_quantity(Decimal('1000'), self.kg), Decimal('1'))

    def test_locations_and_machines_belong_to_same_factory_warehouse(self):
        wh = m.Warehouse.objects.create(code='W1', name='库一', factory=self.factory, kind='raw')
        other = m.Warehouse.objects.create(code='W2', name='库二', factory=self.factory, kind='raw')
        zone = m.StorageLocation.objects.create(warehouse=wh, code='Z1', kind='zone')
        self.create('StorageLocation', {'warehouse': other.pk, 'code': 'B1', 'kind': 'bin', 'parent': zone.pk}, 400)
        self.create('StorageLocation', {'warehouse': wh.pk, 'code': 'B1', 'kind': 'bin', 'parent': zone.pk})
        self.patch('StorageLocation', zone.pk, {'kind': 'bin'}, 400)
        factory = m.Factory.objects.create(code='F2', name='厂二')
        center = m.WorkCenter.objects.create(factory=factory, code='C1', name='中心')
        self.create('Machine', {'code': 'MC1', 'name': '机台', 'factory': self.factory.pk, 'work_center': center.pk}, 400)
        for code in ('MC1', 'MC2'):
            self.assertIsNone(self.create('Machine', {'code': code, 'name': '机台', 'factory': self.factory.pk, 'number': ''})['number'])

    def test_route_approval_freezes_edits_import_bulk_and_deletion(self):
        version = self.version()
        self.assertEqual(self.client.post(self.url('RouteVersion', version.pk, 'approve'), {}, format='json').status_code, 400)
        step = self.step(version)
        self.assertEqual(self.client.post(self.url('RouteVersion', version.pk, 'approve'), {}, format='json').status_code, 200)
        self.patch('RouteVersion', version.pk, {'remarks': '改写'}, 400)
        self.patch('RouteStep', step.pk, {'sequence': 2}, 400)
        self.create('RouteStep', {'route_version': version.pk, 'sequence': 2, 'process': self.process.pk, 'is_required': True}, 400)
        self.assertEqual(self.client.put(self.url('RouteStep', action='batch_update'), [{'id': step.pk, 'sequence': 2}], format='json').status_code, 400)
        self.assertEqual(self.client.delete(self.url('RouteStep', step.pk)).status_code, 400)
        self.assertEqual(self.client.delete(self.url('RouteVersion', action='multiple_delete'), {'keys': [version.pk]}, format='json').status_code, 400)
        response = self.client.get(self.url('RouteStep', action='update_template'), {'id': step.pk})
        book = load_workbook(BytesIO(b''.join(response.streaming_content)))
        for cell in book['数据'][1]:
            if cell.value == '工序号':
                book['数据'].cell(2, cell.column).value = 2
        stream = BytesIO(); book.save(stream); book.close()
        result = self.client.post(self.url('RouteStep', action='import'), {'file': SimpleUploadedFile('节点.xlsx', stream.getvalue())}, format='multipart')
        self.assertEqual(result.status_code, 400, result.data)
        self.assertEqual(m.RouteStep.objects.get(pk=step.pk).sequence, 1)

    def test_route_status_sequence_and_clone_predecessors(self):
        version = self.version()
        first = self.step(version)
        self.step(version, 2, first)
        self.patch('RouteStep', first.pk, {'sequence': 3}, 400)
        self.assertEqual(self.client.post(self.url('RouteVersion', version.pk, 'activate'), {}, format='json').status_code, 400)
        for action in ('approve', 'activate', 'retire'):
            response = self.client.post(self.url('RouteVersion', version.pk, action), {}, format='json')
            self.assertEqual(response.status_code, 200, response.data)
        response = self.client.post(self.url('RouteVersion', version.pk, 'clone'), {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        clone = m.RouteVersion.objects.get(pk=response.data['data']['id'])
        self.assertEqual(clone.version, 2)
        self.assertEqual(clone.status, VersionStatus.DRAFT)
        steps = list(m.RouteStep.objects.filter(route_version=clone).order_by('sequence'))
        self.assertEqual(steps[1].predecessor_id, steps[0].pk)
        self.assertTrue(all(step.legacy_id is None for step in steps))
        self.assertIsNone(clone.legacy_id)
        self.assertEqual(self.client.get(self.url('RouteStep', steps[1].pk)).data['data']['predecessor_label'], '1')

    def test_nullable_active_flag_metadata_preserves_unknown(self):
        fields = {field['key']: field for field in self.client.get('/api/meta/erp_base_party/').data['data']['fields']}
        self.assertTrue(fields['is_active']['nullable'])
        self.assertEqual(fields['is_active']['kind'], 'boolean')
        self.assertIsNone(fields['is_active']['dictionary'])

    def test_manager_viewer_permissions_and_reinitialization(self):
        viewer = Users.objects.create_user(username='viewer')
        viewer.role.add(Role.objects.get(key='erp_base_viewer'))
        self.client.force_authenticate(viewer)
        self.assertEqual(self.client.get(self.url('Party')).status_code, 200)
        self.assertEqual(self.client.post(self.url('Party'), {'code': 'P2', 'name': '无权新增'}, format='json').status_code, 403)
        manager = Users.objects.create_user(username='manager')
        role = Role.objects.get(key='erp_base_manager'); manager.role.add(role)
        self.client.force_authenticate(manager)
        self.create('Party', {'code': 'P2', 'name': '授权新增'})
        FieldPermission.objects.filter(role=role, field__model=m.Party._meta.label, field__field_name='name').update(is_update=False)
        call_command('init_data', verbosity=0)
        self.patch('Party', self.party.pk, {'name': '不可修改'}, 400)
        self.assertEqual(Menu.objects.filter(parent__component_name='erp_base', resource__startswith='erp_base_', visible=True).count(), 20)
        self.assertFalse(RoleMenuPermission.objects.filter(role=role,menu__component_name__startswith='erp_base_legacy').exists())
        self.assertEqual(Menu.objects.filter(parent__component_name='erp_base', resource__startswith='erp_base_').count(), 30)

    def test_related_field_data_scope_prevents_cross_department_link(self):
        dept = Dept.objects.create(key='restricted', name='受限部门')
        actor = Users.objects.create_user(username='scoped', dept=dept)
        role = Role.objects.get(key='erp_base_manager'); actor.role.add(role)
        RoleMenuButtonPermission.objects.filter(role=role, menu_button__value='erp_base_party:list').update(data_range=2)
        self.client.force_authenticate(actor)
        self.create('PartyContact', {'party': self.party.pk, 'name': '越部门联系人'}, 400)
        parent = m.Party.objects.create(code='P2', name='本部门', dept_belong=dept)
        contact = self.create('PartyContact', {'party': parent.pk, 'name': '允许'})
        self.assertEqual(contact['dept_belong'], dept.pk)
