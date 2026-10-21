import csv
import json
from tempfile import TemporaryFile
from types import SimpleNamespace
from zipfile import BadZipFile

from django.core.exceptions import ObjectDoesNotExist, ValidationError as ModelValidationError
from django.db import transaction
from django.http import FileResponse, HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from rest_framework import serializers
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied
from drf_spectacular.utils import extend_schema, inline_serializer

from core.api import success
from core.permissions import ApiPermission, scoped_queryset
from .excel_fields import UPDATE_KEY, SERIAL, KEEP, NULL, EMPTY, columns_for, supports_import, encode, decode

XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
FORMATS = {'text': '@', 'date': 'yyyy-mm-dd', 'datetime': 'yyyy-mm-dd hh:mm:ss', 'time': 'hh:mm:ss'}
GUIDES = {'text': '文本，编号保留前导零', 'number': '数值，0 有效', 'choice': '选择中文选项',
          'relation': '选择名称 [主键]', 'multi': '分号分隔名称 [主键]；含分号的选项用双引号引用',
          'date': 'YYYY-MM-DD', 'datetime': 'YYYY-MM-DD HH:MM:SS', 'time': 'HH:MM:SS', 'json': 'JSON 文本'}


def write_workbook(rows, columns, output, mode='export', context=None):
    book = Workbook()
    sheet = book.active
    sheet.title = '数据'
    sheet.freeze_panes = 'A2'
    guide = book.create_sheet('填写说明')
    guide.append(['项目', '说明'])
    guide.append(['用途', {'export': '数据导出', 'create': '新增模板', 'update': '批量更新模板'}[mode]])
    guide.append(['时间', f'日期时间采用项目时区 {timezone.get_current_timezone_name()}'])
    guide.append(['空白/保留', '空单元格或 #KEEP 不提交；新增采用默认值，更新保持原值。0、False、负数和空格文本均为有效值。'])
    guide.append(['明确清空', '#NULL 表示数据库NULL；#EMPTY 表示空字符串（多选字段表示清空选项）。字段不接受该值时报告错误，不转换为0。'])
    guide.append(['标记原文', '原文本恰好为 #NULL、#EMPTY、#KEEP 或以 #TEXT: 开头时，在前面加 #TEXT:，例如 #TEXT:#NULL。'])
    guide.append(['公式', '本系统不计算Excel公式；只读取Excel/WPS保存的计算结果，公式无缓存或结果错误时报告具体位置，不作为空值跳过。'])
    guide.append(['数据行', '完全空白行仅计为空白；有序号的行明确参与导入。重复行逐行处理；任何一行失败则整批回滚。'])
    if context:
        config=book.create_sheet('导入设置');config.append(['设置项','值'])
        for key,value in context.items():config.append([key,value])
        guide.append(['默认业务入口', '数据行“业务入口”优先，其次“导入设置”的业务入口，最后使用导入窗口所选入口。更新已有单据保留原入口。'])
    if mode != 'export':
        guide.append(['文件', '表格不提交文件内容；更新保留已有附件，新增所需的必填文件须通过上传表单提交'])
    if mode == 'update':
        guide.append(['更新主键', '用于定位已有记录，不能删掉该列；不转为新增'])
    guide.append(['字段', '填写方式'])
    for column in columns:
        guide.append([column['title'], ('新增必填；' if column['required'] else '') + GUIDES[column['kind']] + ('；'+column['guide'] if 'guide' in column else '')])
    guide.column_dimensions['A'].width = 24
    guide.column_dimensions['B'].width = 90
    sheet.append([SERIAL] + [column['title'] for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2563EB')
    for index, column in enumerate(columns, 2):
        sheet.cell(1, index).comment = Comment(GUIDES[column['kind']], '管理平台')
        sheet.column_dimensions[get_column_letter(index)].width = max(18, len(column['title']) * 2 + 4)
        sheet.column_dimensions[get_column_letter(index)].number_format = FORMATS.get(column['kind'], 'General')
    sheet.column_dimensions['A'].width = 10
    if mode != 'export' and any(column['options'] for column in columns):
        options = book.create_sheet('选项')
        options.sheet_state = 'hidden'
        for index, column in enumerate(columns, 1):
            if not column['options']:
                continue
            letter = get_column_letter(index)
            options.cell(1, index, column['title'])
            for row, option in enumerate(column['options'], 2):
                options.cell(row, index, option['label']).data_type = 's'
            name = f'Options_{index}'
            book.defined_names.add(DefinedName(name, attr_text=
                f'{quote_sheetname(options.title)}!${letter}$2:${letter}${len(column["options"])+1}'))
            if column['kind'] != 'multi':
                validation = DataValidation(type='list', formula1='='+name, allow_blank=True,
                                            showDropDown=False, showInputMessage=True, prompt='请选择选项')
                sheet.add_data_validation(validation)
                validation.add(f'{get_column_letter(index+1)}2:{get_column_letter(index+1)}1048576')
            else:
                guide.append([column['title']+' 示例', encode([option['value'] for option in column['options'][:2]], column)])
    count = 0
    for count, row in enumerate(rows, 1):
        sheet.append([count] + [encode(row.get(column['key']), column, for_import=mode!='export') for column in columns])
        for index, column in enumerate(columns, 2):
            cell = sheet.cell(count+1, index)
            if cell.value is not None:
                cell.number_format = FORMATS.get(column['kind'], 'General')
                if column['kind'] == 'number' and column['decimals'] is not None:
                    cell.number_format = '0' + ('.'+'0'*column['decimals'] if column['decimals'] else '')
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    if not count:
        sheet.append([None] * (len(columns)+1))
        for index, column in enumerate(columns, 2):
            sheet.cell(2, index).number_format = FORMATS.get(column['kind'], 'General')
    table = Table(displayName='DataRows', ref=f'A1:{get_column_letter(len(columns)+1)}{max(count+1, 2)}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleLight11', showRowStripes=True)
    sheet.add_table(table)
    book.save(output)
    book.close()


def excel_response(rows, columns, filename, mode='export', context=None):
    output = TemporaryFile()
    try:
        write_workbook(rows, columns, output, mode, context)
        output.seek(0)
        return FileResponse(output, as_attachment=True, filename=filename+'.xlsx', content_type=XLSX_TYPE)
    except Exception:
        output.close()
        raise


class ExcelActions:
    def excel_context(self, book=None):
        return {}

    def excel_prepare(self, data, instance, context):
        return data

    def excel_serializer(self,instance,data,mode,context):
        return self.get_serializer(instance,data=data,partial=mode=='update')

    def excel_finish(self,context,counts):
        pass

    def excel_save(self, serializer, audit, mode):
        return serializer.save(**audit)

    def excel_query_request(self, method):
        base = self.request.path.rsplit('/', 2)[0]+'/'
        return SimpleNamespace(user=self.request.user, method=method, path=base+('0/' if method == 'PATCH' else ''))

    def excel_write_queryset(self):
        request = self.excel_query_request('PATCH')
        if not ApiPermission().has_permission(request, self):
            raise PermissionDenied('无修改权限')
        return scoped_queryset(self.queryset.all(), request)

    def excel_rows(self, queryset, columns, update=False):
        relations = [column['key'] for column in columns if column['kind'] in ['multi', 'relation']]
        if relations:
            queryset = queryset.prefetch_related(*relations)
        if any(column.get('source') == 'dept.owner' for column in columns):
            queryset = queryset.select_related('dept')
        serializer = self.get_serializer()
        for obj in queryset.iterator(chunk_size=2000):
            row = dict(serializer.to_representation(obj))
            for column in columns:
                if column.get('source') == 'dept.owner':
                    row[column['key']] = obj.dept.owner if obj.dept else ''
            if update:
                row['__pk'] = obj.pk
            yield row

    def excel_columns(self, mode, queryset=None):
        columns = columns_for(self, mode, queryset)
        if mode == 'update':
            columns.insert(0, {'key': '__pk', 'title': UPDATE_KEY, 'kind': 'text', 'required': True, 'options': []})
        return columns

    @action(methods=['get'], detail=False)
    def export(self, request):
        queryset = self.filter_queryset(self.get_queryset())
        columns = self.excel_columns('export', queryset)
        return excel_response(self.excel_rows(queryset, columns), columns, str(self.queryset.model._meta.verbose_name)+'导出')

    @action(methods=['get'], detail=False)
    def import_template(self, request):
        if not supports_import(self.queryset.model):
            raise serializers.ValidationError('该模块只提供数据导出')
        return excel_response([], self.excel_columns('create'), str(self.queryset.model._meta.verbose_name)+'新增模板', 'create', self.excel_context())

    @action(methods=['get'], detail=False)
    def update_template(self, request):
        if not supports_import(self.queryset.model):
            raise serializers.ValidationError('该模块只提供数据导出')
        queryset = self.filter_queryset(self.get_queryset()).filter(pk__in=self.excel_write_queryset().values('pk'))
        columns = self.excel_columns('update', queryset)
        return excel_response(self.excel_rows(queryset, columns, update=True), columns,
                              str(self.queryset.model._meta.verbose_name)+'更新模板', 'update', self.excel_context())

    @extend_schema(request=inline_serializer(name='ExcelImportFile', fields={'file': serializers.FileField()}))
    @action(methods=['post'], detail=False, url_path='import')
    @transaction.atomic
    def import_data(self, request):
        if not supports_import(self.queryset.model):
            raise serializers.ValidationError('该模块只提供数据导出')
        upload = request.FILES['file']
        if not upload.name.lower().endswith('.xlsx'):
            raise serializers.ValidationError('请上传 xlsx 文件')
        try:
            book = load_workbook(upload, read_only=True, data_only=False)
            upload.seek(0)
            cached = load_workbook(upload, read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile) as exc:
            raise serializers.ValidationError('文件不是有效的 xlsx 工作簿') from exc
        try:
            if '数据' not in book.sheetnames:
                raise serializers.ValidationError('模板缺少“数据”工作表')
            sheet = book['数据']
            rows = sheet.iter_rows()
            headers = tuple(cell.value for cell in next(rows, ()))
            cached_rows=cached['数据'].iter_rows(min_row=2)
            if not any(header and header != SERIAL for header in headers):
                raise serializers.ValidationError('模板缺少中文表头')
            mode = 'update' if UPDATE_KEY in headers else 'create'
            queryset = self.excel_write_queryset() if mode == 'update' else None
            if mode == 'create' and not ApiPermission().has_permission(self.excel_query_request('POST'), self):
                raise PermissionDenied('无新增权限')
            context=self.excel_context(book)
            columns = {column['title']: column for column in self.excel_columns(mode, queryset)}
            titles = [header for header in headers if header and header != SERIAL]
            if len(titles) != len(set(titles)):
                raise serializers.ValidationError({'sheet': '数据', 'row': 1, 'field': '表头', 'detail': '表头重复'})
            for title in titles:
                if title not in columns:
                    raise serializers.ValidationError({'sheet': '数据', 'row': 1, 'field': str(title), 'detail': '字段不在本模板可导入范围'})
            counts = {'created': 0, 'updated': 0, 'skipped': 0}
            row_results=[]
            for number, (cells,cached_cells) in enumerate(zip(rows,cached_rows), 2):
                data, primary = {}, None
                explicit_fields=set()
                title = '更新主键' if mode == 'update' else '数据'
                try:
                    if all(cell.value is None or cell.value=='' for cell in cells):
                        counts['skipped']+=1
                        row_results.append({'row':number,'status':'blank'})
                        continue
                    for title, cell, cached_cell in zip(headers, cells,cached_cells):
                        value=cell.value
                        if cell.data_type=='f':
                            value=cached_cell.value
                            if value is None:raise serializers.ValidationError('Excel公式没有已保存的计算结果，请计算并保存后再导入。')
                        if cell.data_type=='e' or cached_cell.data_type=='e':raise serializers.ValidationError('Excel单元格结果为错误：'+str(value))
                        if title in [None,''] and value not in [None,'']:raise serializers.ValidationError('有内容的列未填写表头：'+cell.column_letter)
                        if title in [None, SERIAL] or value is None or value == '' or value == KEEP:
                            continue
                        column = columns[title]
                        if column['key'] == '__pk':
                            primary = decode(value,column)
                        else:
                            data[column['key']] = decode(value, column)
                            if value in [NULL,EMPTY]:explicit_fields.add(column['key'])
                    instance = queryset.get(pk=primary) if mode == 'update' else None
                    data=self.excel_prepare(data,instance,context)
                    serializer = self.excel_serializer(instance,data,mode,context)
                    serializer.context['excel_explicit_fields']=explicit_fields
                    serializer.context['excel_row']=number
                    serializer.is_valid(raise_exception=True)
                    audit = self.audit()
                    if mode == 'update':
                        audit.pop('creator', None)
                        audit.pop('dept_belong_id', None)
                    saved=self.excel_save(serializer, audit, mode)
                    counts['updated' if mode == 'update' else 'created'] += 1
                    row_results.append({'row':number,'status':'updated' if mode=='update' else 'created','id':saved.pk})
                except (KeyError, IndexError, ValueError, TypeError, csv.Error, ObjectDoesNotExist, serializers.ValidationError,ModelValidationError) as exc:
                    if isinstance(exc, ObjectDoesNotExist):
                        title, detail = UPDATE_KEY, '记录不存在或不在修改范围'
                    elif isinstance(exc, (serializers.ValidationError,ModelValidationError)):
                        detail = exc.detail if isinstance(exc,serializers.ValidationError) else serializers.as_serializer_error(exc)
                        if isinstance(detail,dict):
                            title = '、'.join(c['title'] for c in columns.values() if c['key'] in detail) or '数据'
                    else:
                        detail = '选项或填写格式不正确：'+str(exc)
                    raise serializers.ValidationError({'sheet': '数据', 'row': number, 'field': title, 'detail': detail}) from exc
            self.excel_finish(context,counts)
            return success({'mode': mode, 'count': counts['created']+counts['updated'], **counts,'rows':row_results}, '导入成功')
        finally:
            cached.close()
            book.close()

    @action(methods=['post'], detail=False)
    def async_export(self, request):
        from .models import DownloadCenter
        from .tasks import export_file
        from application.celery import app
        queryset = self.filter_queryset(self.get_queryset())
        columns = self.excel_columns('export', queryset)
        rows = list(self.excel_rows(queryset, columns))
        item = DownloadCenter.objects.create(task_name=str(self.queryset.model._meta.verbose_name)+'导出', **self.audit())
        result = app.send_task(export_file.name, args=[item.pk, rows, columns])
        item.task_id = result.id
        item.save(update_fields=['task_id'])
        return success({'id': item.pk}, '导出任务已提交')

    @action(methods=['get'], detail=False)
    def init_data(self, request):
        rows = self.get_serializer(self.filter_queryset(self.get_queryset()), many=True).data
        response = HttpResponse(json.dumps(list(rows), ensure_ascii=False, indent=2), content_type='application/json')
        response['Content-Disposition'] = 'attachment; filename="init-data.json"'
        return response
