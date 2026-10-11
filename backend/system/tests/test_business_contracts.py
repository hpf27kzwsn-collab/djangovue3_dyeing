from io import BytesIO
from tempfile import TemporaryDirectory
from unittest import skipUnless

from django.apps import apps
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import connection
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from system.models import Users, Role, FieldPermission, MenuField, MenuButton, RoleMenuButtonPermission, FileList


@skipUnless(apps.is_installed('system.tests.twin_a'), 'requires extension_settings')
class BusinessContractTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('init_data', verbosity=0)

    def setUp(self):
        self.client = APIClient()
        self.admin = Users.objects.get(username='superadmin')
        self.client.force_authenticate(self.admin)

    def workbook(self, entity, action, **query):
        response = self.client.get(f'/api/system/{entity}/{action}/', query)
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        result = load_workbook(BytesIO(b''.join(response.streaming_content)))
        return result

    def upload(self, entity, book):
        content = BytesIO()
        book.save(content)
        book.close()
        return self.client.post(f'/api/system/{entity}/import/',
            {'file': SimpleUploadedFile('业务.xlsx', content.getvalue())}, format='multipart')

    def fill(self, book, values, row=2):
        for cell in book['数据'][1]:
            if cell.value in values:
                book['数据'].cell(row, cell.column).value = values[cell.value]

    def test_business_authentication_names_defaults_edit_and_xlsx(self):
        from system.tests.twin_a.models import BusinessCredential
        entity = 'identity_a_businesscredential'
        meta = self.client.get('/api/meta/'+entity+'/').data['data']
        fields = {f['key']: f for f in meta['fields']}
        self.assertEqual(fields['password']['default'], 'device-initial')
        self.assertEqual(fields['password']['kind'], 'text')
        self.assertTrue(fields['password']['is_update'])
        self.assertEqual(set(MenuField.objects.filter(model=BusinessCredential._meta.label).values_list('field_name', flat=True)),
                         {'description', 'name', 'password', 'groups', 'user_permissions'})
        item = BusinessCredential.objects.create(name='设备')
        response = self.client.patch(f'/api/system/{entity}/{item.pk}/', {'password':'device-new'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        item.refresh_from_db();self.assertEqual(item.password, 'device-new')
        table = next(t for t in self.client.get('/api/system/system_config/association_tables/').data['data'] if t['value']==entity)
        self.assertTrue({'password', 'groups', 'user_permissions'} <= {f['key'] for f in table['fields']})
        book = self.workbook(entity, 'import_template')
        self.fill(book, {'设备名称':'xlsx设备', '设备口令':'xlsx-old', '业务分组':'A', '业务许可说明':'许可'})
        self.assertEqual(self.upload(entity, book).status_code, 200)
        imported = BusinessCredential.objects.get(name='xlsx设备')
        book = self.workbook(entity, 'update_template', id=imported.pk)
        self.fill(book, {'设备口令':'xlsx-new'})
        self.assertEqual(self.upload(entity, book).status_code, 200)
        imported.refresh_from_db();self.assertEqual(imported.password, 'xlsx-new')
        exported = self.workbook(entity, 'export')
        self.assertIn('设备口令', [c.value for c in exported['数据'][1]])
        exported.close()
        internal = self.client.get('/api/system/system_config/association_tables/').data['data']
        self.assertNotIn('password', [f['key'] for t in internal if t['value']=='user' for f in t['fields']])
        Users.objects.create_user(username='internal-other', password='internal-fixture')
        self.assertNotIn('password', self.client.get('/api/system/user/').data['data'][0])

    def test_business_authentication_names_query_create_update_grants(self):
        from system.tests.twin_a.models import BusinessCredential
        entity = 'identity_a_businesscredential'
        role = Role.objects.create(name='设备角色', key='business-contract')
        actor = Users.objects.create_user(username='business-contract')
        actor.role.add(role)
        for action in ['list', 'add', 'patch', 'export-get', 'import-template-get']:
            RoleMenuButtonPermission.objects.create(role=role, menu_button=MenuButton.objects.get(value=f'{entity}:{action}'), data_range=3)
        item = BusinessCredential.objects.create(name='授权设备', password='secret')
        self.client.force_authenticate(actor)
        row = self.client.get(f'/api/system/{entity}/').data['data'][0]
        self.assertFalse({'password', 'groups', 'user_permissions'} & row.keys())
        self.assertEqual(self.client.patch(f'/api/system/{entity}/{item.pk}/', {'password':'blocked'}, format='json').status_code, 400)
        grants = {}
        for key in ['name', 'password', 'groups', 'user_permissions']:
            field = MenuField.objects.get(model=BusinessCredential._meta.label, field_name=key)
            grants[key] = FieldPermission.objects.create(role=role, field=field, is_query=True, is_create=True, is_update=False)
        self.assertEqual(self.client.get(f'/api/system/{entity}/').data['data'][0]['password'], 'secret')
        created = self.client.post(f'/api/system/{entity}/', {'name':'授权新增','password':'created','groups':'B','user_permissions':'C'}, format='json')
        self.assertEqual(created.status_code, 200, created.data)
        self.assertEqual(self.client.patch(f'/api/system/{entity}/{item.pk}/', {'password':'blocked'}, format='json').status_code, 400)
        grants['password'].is_update=True;grants['password'].save()
        self.assertEqual(self.client.patch(f'/api/system/{entity}/{item.pk}/', {'password':'allowed'}, format='json').status_code, 200)
        item.refresh_from_db();self.assertEqual(item.password, 'allowed')

    def test_business_key_rename_preserves_self_fk_m2m_and_alternate_target(self):
        from system.tests.twin_a.models import BusinessKey, KeyReference, Twin
        tag = Twin.objects.create(name='标签')
        item = BusinessKey.objects.create(id='OLD', name='主记录', reference='REF')
        item.tags.add(tag)
        child = BusinessKey.objects.create(id='CHILD', name='子记录', reference='CHILD-REF', parent=item)
        external = KeyReference.objects.create(target=item, alternate=item)
        external.targets.add(item)
        response = self.client.patch('/api/system/identity_a_businesskey/OLD/', {'id':'NEW','name':'改名','parent':'OLD','tags':[tag.pk]}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['data']['id'], 'NEW')
        self.assertFalse(BusinessKey.objects.filter(pk='OLD').exists())
        self.assertEqual(BusinessKey.objects.count(), 2)
        renamed = BusinessKey.objects.get(pk='NEW')
        self.assertEqual((renamed.parent_id,renamed.name), ('NEW','改名'))
        child.refresh_from_db();external.refresh_from_db()
        self.assertEqual(child.parent_id, 'NEW');self.assertEqual(external.target_id, 'NEW')
        self.assertEqual(external.alternate_id, 'REF')
        self.assertEqual(list(external.targets.values_list('pk',flat=True)), ['NEW'])
        self.assertEqual(list(renamed.tags.values_list('pk',flat=True)), [tag.pk])
        connection.check_constraints()

    def test_business_key_rename_follows_primary_foreign_key_chain(self):
        from system.tests.twin_a.models import BusinessKey, KeyIdentity, KeyIdentityReference
        item = BusinessKey.objects.create(id='OLD', name='主记录', reference='REF')
        identity = KeyIdentity.objects.create(target=item)
        reference = KeyIdentityReference.objects.create(target=identity)
        result = self.client.patch('/api/system/identity_a_businesskey/OLD/', {'id':'NEW'}, format='json')
        self.assertEqual(result.status_code, 200, result.data)
        self.assertTrue(KeyIdentity.objects.filter(pk='NEW').exists())
        reference.refresh_from_db();self.assertEqual(reference.target_id, 'NEW')
        connection.check_constraints()

    def test_business_key_collision_does_not_change_existing_record(self):
        from system.tests.twin_a.models import BusinessKey
        BusinessKey.objects.create(id='OLD', name='原记录', reference='REF')
        BusinessKey.objects.create(id='TAKEN', name='其他记录', reference='OTHER')
        result = self.client.patch('/api/system/identity_a_businesskey/OLD/', {'id':'TAKEN','name':'不应保存'}, format='json')
        self.assertEqual(result.status_code, 400)
        self.assertEqual(BusinessKey.objects.get(pk='OLD').name, '原记录')
        self.assertEqual(BusinessKey.objects.count(), 2)
        meta = self.client.get('/api/meta/identity_a_businesskey/').data['data']
        self.assertTrue(next(f for f in meta['fields'] if f['key']=='id')['is_update'])
        self.assertTrue(MenuField.objects.filter(model=BusinessKey._meta.label, field_name='id').exists())
        self.assertTrue(next(f for f in self.client.get('/api/meta/post/').data['data']['fields'] if f['key']=='id')['readonly'])
        role = Role.objects.create(name='主键授权角色', key='key-grant')
        actor = Users.objects.create_user(username='key-grant')
        actor.role.add(role)
        RoleMenuButtonPermission.objects.create(role=role,
            menu_button=MenuButton.objects.get(value='identity_a_businesskey:patch'), data_range=3)
        field = MenuField.objects.get(model=BusinessKey._meta.label, field_name='id')
        grant = FieldPermission.objects.create(role=role, field=field, is_query=True, is_create=False, is_update=False)
        self.client.force_authenticate(actor)
        self.assertEqual(self.client.patch('/api/system/identity_a_businesskey/OLD/', {'id':'GRANTED'}, format='json').status_code, 400)
        grant.is_update=True;grant.save()
        result = self.client.patch('/api/system/identity_a_businesskey/OLD/', {'id':'GRANTED'}, format='json')
        self.assertEqual(result.status_code, 200, result.data)
        self.assertTrue(BusinessKey.objects.filter(pk='GRANTED').exists())
        self.assertEqual(BusinessKey.objects.count(), 2)

    def test_optional_and_excluded_files_do_not_disable_xlsx_data(self):
        from system.tests.twin_a.models import FileEntry, ExcludedFileEntry
        for model in [FileEntry, ExcludedFileEntry]:
            with self.subTest(model=model.__name__):
                entity = f'identity_a_{model._meta.model_name}'
                self.assertTrue(self.client.get('/api/meta/'+entity+'/').data['data']['excel_import'])
                book = self.workbook(entity, 'import_template')
                self.assertNotIn('可选附件', [c.value for c in book['数据'][1]])
                self.fill(book, {'资料名称':'xlsx文字'})
                self.assertEqual(self.upload(entity, book).status_code, 200)
                item = model.objects.get(name='xlsx文字')
                model.objects.filter(pk=item.pk).update(attachment='fixture/existing.pdf')
                book = self.workbook(entity, 'update_template', id=item.pk)
                self.fill(book, {'资料名称':'xlsx修改'})
                self.assertEqual(self.upload(entity, book).status_code, 200)
                item.refresh_from_db()
                self.assertEqual((item.name,str(item.attachment)), ('xlsx修改','fixture/existing.pdf'))

    def test_required_file_reports_real_field_error_instead_of_module_ban(self):
        entity = 'identity_a_requiredfileentry'
        book = self.workbook(entity, 'import_template')
        self.fill(book, {'资料名称':'缺少必填附件'})
        result = self.upload(entity, book)
        self.assertEqual(result.status_code, 400)
        self.assertIn('attachment', str(result.data))
        self.assertNotIn('只提供数据导出', str(result.data))

    def test_filelist_xlsx_updates_metadata_preserving_uploaded_file(self):
        with TemporaryDirectory() as directory, override_settings(MEDIA_ROOT=directory):
            original = b'%PDF-1.4\nfixture content\n%%EOF'
            response = self.client.post('/api/system/file/',
                {'url': SimpleUploadedFile('existing.pdf', original, content_type='application/pdf')}, format='multipart')
            self.assertEqual(response.status_code, 200, response.data)
            item = FileList.objects.get(pk=response.data['data']['id'])
            original_path = str(item.url)
            book = self.workbook('file', 'update_template', id=item.pk)
            self.assertNotIn('文件', [c.value for c in book['数据'][1]])
            title = FileList._meta.get_field('name').verbose_name
            self.fill(book, {str(title):'xlsx附件名'})
            result = self.upload('file', book)
            self.assertEqual(result.status_code, 200, result.data)
            item.refresh_from_db()
            self.assertEqual((item.name,str(item.url)), ('xlsx附件名',original_path))
            with item.url.open('rb') as stream:
                self.assertEqual(stream.read(), original)
    def test_large_page_prefetch_respects_sqlite_parameters_without_truncation(self):
        from system.tests.twin_a.models import PlainEntry, Twin
        tag = Twin.objects.create(name='large-page-tag')
        rows = PlainEntry.objects.bulk_create([PlainEntry(name=f'large-{i}') for i in range(33000)])
        rows[0].tags.add(tag)
        rows[-1].tags.add(tag)
        response = self.client.get('/api/system/identity_a_plainentry/', {'limit':33000,'ordering':'id'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.data['total'],33000)
        self.assertEqual(len(response.data['data']),33000)
        self.assertEqual(response.data['data'][0]['tags'],[tag.pk])
        self.assertEqual(response.data['data'][-1]['tags'],[tag.pk])
