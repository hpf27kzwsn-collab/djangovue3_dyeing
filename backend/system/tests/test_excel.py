import csv
from datetime import datetime
from hashlib import md5
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch

from django.core.management import call_command
from django.test import TestCase, override_settings
from django.utils import timezone
from django_celery_beat.models import ClockedSchedule
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient

from system.models import Users, Dept, Role, Post, Area, SystemConfig, MenuField, MenuButton, FieldPermission, RoleMenuButtonPermission, DownloadCenter, MessageCenter, MessageCenterTargetUser
from system.tasks import export_file
from system.excel_fields import UPDATE_KEY


class ExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('init_data', verbosity=0)

    def setUp(self):
        self.client = APIClient()
        self.root = Users.objects.get(username='superadmin')
        self.client.force_authenticate(self.root)

    def template(self, entity, update=False, **filters):
        response = self.client.get(f'/api/system/{entity}/{"update_template" if update else "import_template"}/', filters)
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        content = b''.join(response.streaming_content)
        return load_workbook(BytesIO(content))

    def fill(self, book, records):
        sheet = book['数据']
        for row, record in enumerate(records, 2):
            for column, header in enumerate(sheet[1], 1):
                sheet.cell(row, column).value = record.get(header.value)

    def upload(self, entity, book, name='数据.xlsx'):
        output = BytesIO()
        book.save(output)
        book.close()
        return self.client.post(f'/api/system/{entity}/import/',
            {'file': SimpleUploadedFile(name, output.getvalue())}, format='multipart')

    def export_book(self, entity, **filters):
        response = self.client.get(f'/api/system/{entity}/export/', filters)
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        content = b''.join(response.streaming_content)
        return load_workbook(BytesIO(content))

    def test_import_message_runs_normal_delivery_and_requires_headers(self):
        recipient = Users.objects.create_user(username='excel-recipient', password='recipient-password')
        book = self.template('message_center')
        self.fill(book, [{'标题': '导入分发', '内容': '消息正文', '消息类型': '指定用户',
                         '接收用户': f'{recipient.username} [{recipient.pk}]'}])
        result = self.upload('message_center', book)
        self.assertEqual(result.status_code, 200, result.data)
        message = MessageCenter.objects.get(title='导入分发')
        self.assertTrue(MessageCenterTargetUser.objects.filter(messagecenter=message, users=recipient).exists())
        book = self.template('post')
        for cell in book['数据'][1]:
            cell.value = None
        self.assertEqual(self.upload('post', book).status_code, 400)

    def test_native_datetime_import_export_update_uses_project_timezone(self):
        title = str(ClockedSchedule._meta.get_field('clocked_time').verbose_name)
        local_time = datetime(2026, 11, 9, 15, 6, 7)
        book = self.template('clocked_schedule')
        self.fill(book, [{title: local_time}])
        result = self.upload('clocked_schedule', book)
        self.assertEqual(result.status_code, 200, result.data)
        item = ClockedSchedule.objects.get()
        self.assertTrue(timezone.is_aware(item.clocked_time))
        self.assertEqual(timezone.localtime(item.clocked_time).replace(tzinfo=None), local_time)
        book = self.template('clocked_schedule', update=True)
        sheet = book['数据']
        column = next(cell.column for cell in sheet[1] if cell.value == title)
        self.assertEqual(sheet.cell(2, column).value, local_time)
        changed = datetime(2026, 11, 10, 9, 8, 7)
        sheet.cell(2, column).value = changed
        result = self.upload('clocked_schedule', book)
        self.assertEqual(result.status_code, 200, result.data)
        item.refresh_from_db()
        self.assertEqual(timezone.localtime(item.clocked_time).replace(tzinfo=None), changed)

    def test_new_user_template_and_readable_relations(self):
        dept = Dept.objects.create(name='财务部')
        first = Role.objects.create(name='同名角色', key='a')
        second = Role.objects.create(name='同名角色', key='b')
        punctuated = Role.objects.create(name='采购;兼任 "审核"', key='c')
        book = self.template('user')
        self.assertEqual(book['选项'].sheet_state, 'hidden')
        self.assertEqual(book['数据'].freeze_panes, 'A2')
        self.assertNotIn('password', [cell.value for cell in book['数据'][1]])
        self.assertTrue(book.defined_names)
        self.assertTrue(book['数据'].data_validations.dataValidation)
        stream = StringIO()
        csv.writer(stream, delimiter=';', lineterminator='').writerow(
            [f'同名角色 [{first.pk}]', f'同名角色 [{second.pk}]', f'{punctuated.name} [{punctuated.pk}]'])
        self.fill(book, [{'登录账号':'0012300', '用户名称':'导入用户', '手机号码':'00123456789',
                          '用户性别':'女', '账号状态':'禁用', '部门':f'财务部 [{dept.pk}]', '角色':stream.getvalue()}])
        response = self.upload('user', book)
        self.assertEqual(response.status_code, 200, response.data)
        user = Users.objects.get(username='0012300')
        self.assertEqual((user.gender, user.is_active, user.mobile), (2, False, '00123456789'))
        self.assertEqual((user.dept_id, user.dept_belong_id), (dept.pk, dept.pk))
        self.assertEqual(set(user.role.values_list('pk', flat=True)), {first.pk, second.pk, punctuated.pk})
        self.assertTrue(user.check_password('admin123456'))

    def test_update_preserves_password_and_handles_blank_zero_and_duplicate_rows(self):
        dept = Dept.objects.create(name='新部门')
        user = Users.objects.create_user(username='moving', password='keep-password', name='原名称', mobile='0012',
                                         dept=self.root.dept, dept_belong=self.root.dept)
        book = self.template('user', update=True, username='moving')
        self.fill(book, [{UPDATE_KEY:user.pk, '用户名称':'改名', '手机号码':None, '用户性别':0,
                          '账号状态':False, '部门':f'新部门 [{dept.pk}]'},
                         {UPDATE_KEY:user.pk, '用户名称':'最终名称'}])
        response = self.upload('user', book)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['data']['updated'], 2)
        user.refresh_from_db()
        self.assertEqual((user.name, user.mobile, user.gender, user.is_active), ('最终名称', '0012', 0, False))
        self.assertEqual((user.dept_id, user.dept_belong_id), (dept.pk, dept.pk))
        self.assertTrue(user.check_password('keep-password'))
        self.assertEqual(Users.objects.filter(username='moving').count(), 1)

    def test_update_unknown_primary_rolls_back_prior_rows(self):
        post = Post.objects.create(name='原岗位', code='p')
        book = self.template('post', update=True)
        self.fill(book, [{UPDATE_KEY:post.pk, '岗位名称':'不应保存'}, {UPDATE_KEY:999999, '岗位名称':'不应新增'}])
        response = self.upload('post', book)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(str(response.data['data']['row']), '3')
        self.assertEqual(str(response.data['data']['field']), UPDATE_KEY)
        post.refresh_from_db()
        self.assertEqual(post.name, '原岗位')
        self.assertEqual(Post.objects.count(), 1)

    def test_create_error_rolls_back_and_column_order_is_by_title(self):
        book = self.template('role')
        self.fill(book, [{'角色名称':'第一条', '权限字符':'duplicate'}, {'角色名称':'第二条', '权限字符':'duplicate'}])
        response = self.upload('role', book)
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Role.objects.filter(key='duplicate').exists())
        book = self.template('post')
        sheet = book['数据']
        sheet.delete_rows(1, sheet.max_row)
        sheet.append(['岗位编码', '岗位状态', '岗位名称', '排序'])
        sheet.append(['0007', '离职', '重排岗位', 0])
        response = self.upload('post', book)
        self.assertEqual(response.status_code, 200, response.data)
        post = Post.objects.get(code='0007')
        self.assertEqual((post.name, post.status, post.sort), ('重排岗位', 0, 0))

    def test_field_and_update_data_scope_and_not_upsert(self):
        dept = Dept.objects.create(name='本部门')
        other = Dept.objects.create(name='外部门')
        user = Users.objects.create_user(username='excel-staff', dept=dept)
        role = Role.objects.create(name='表格权限', key='excel-staff'); user.role.add(role)
        own = Post.objects.create(name='自己的', code='keep', dept_belong=dept)
        foreign = Post.objects.create(name='外部', code='foreign', dept_belong=other)
        for value in ['post:update-template-get', 'post:import-data-post', 'post:patch']:
            RoleMenuButtonPermission.objects.create(role=role, menu_button=MenuButton.objects.get(value=value), data_range=2)
        for field in MenuField.objects.filter(model='system.Post'):
            FieldPermission.objects.create(role=role, field=field, is_query=True, is_update=field.field_name!='code')
        self.client.force_authenticate(user)
        book = self.template('post', update=True)
        headers = [cell.value for cell in book['数据'][1]]
        self.assertNotIn('岗位编码', headers)
        ids = [row[1] for row in book['数据'].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(ids, [str(own.pk)])
        self.fill(book, [{UPDATE_KEY:foreign.pk, '岗位名称':'越范围'}])
        response = self.upload('post', book)
        self.assertEqual(response.status_code, 400)
        foreign.refresh_from_db(); self.assertEqual(foreign.name, '外部')
        book = self.template('post', update=True)
        sheet = book['数据']; sheet.cell(1, sheet.max_column+1, '岗位编码')
        response = self.upload('post', book)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(str(response.data['data']['row']), '1')
        own.refresh_from_db(); self.assertEqual(own.code, 'keep')

    def test_string_primary_json_dates_and_literal_text(self):
        area = Area.objects.create(code='00101', name='原地区')
        book = self.template('area', update=True)
        self.fill(book, [{UPDATE_KEY:area.pk, '名称':'新地区'}])
        response = self.upload('area', book)
        self.assertEqual(response.status_code, 200, response.data)
        area.refresh_from_db(); self.assertEqual(area.name, '新地区')
        book = self.template('system_config')
        self.fill(book, [{'标题':'=字面内容', '配置键':'roundtrip', '配置值':'{"金额":0,"开关":false}'}])
        # 用户文本以等号开头并非公式。
        sheet = book['数据']
        for cell in sheet[2]:
            if cell.value == '=字面内容': cell.data_type='s'
        response = self.upload('system_config', book)
        self.assertEqual(response.status_code, 200, response.data)
        config = SystemConfig.objects.get(key='roundtrip')
        self.assertEqual(config.value, {'金额':0, '开关':False})
        exported = self.export_book('system_config', key='roundtrip')
        cells = {header.value:cell for header,cell in zip(exported['数据'][1], exported['数据'][2])}
        self.assertEqual(cells['标题'].data_type, 's')
        self.assertEqual(cells['标题'].value, '=字面内容')
        self.assertIsInstance(cells['创建时间'].value, datetime)
        self.assertIsNone(cells['创建时间'].value.tzinfo)
        exported.close()

    def test_xlsx_only_and_read_only_entities(self):
        response = self.upload('post', self.template('post'), name='旧文件.xls')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(self.client.get('/api/meta/operation_log/').data['data']['excel_import'])
        response = self.client.get('/api/system/operation_log/import_template/')
        self.assertEqual(response.status_code, 400)

    def test_explicit_null_and_empty_clear_relations_and_multiple_choices(self):
        role=Role.objects.create(name='待清空角色',key='clear-role')
        user=Users.objects.create_user(username='clear-user',mobile='001',dept=self.root.dept)
        user.role.add(role)
        book=self.template('user',update=True,username='clear-user')
        self.fill(book,[{UPDATE_KEY:user.pk,'手机号码':'#EMPTY','部门':'#NULL','角色':'#EMPTY','账号状态':False,'用户性别':0}])
        response=self.upload('user',book)
        self.assertEqual(response.status_code,200,response.data)
        user.refresh_from_db()
        self.assertEqual((user.mobile,user.dept_id,user.is_active,user.gender),('',None,False,0))
        self.assertFalse(user.role.exists())

    def test_nonnullable_null_is_error_and_rolls_back_all_rows(self):
        book=self.template('post')
        self.fill(book,[{'岗位名称':'有效岗位','岗位编码':'valid','排序':1},{'岗位名称':'空排序','岗位编码':'null','排序':'#NULL'}])
        response=self.upload('post',book)
        self.assertEqual(response.status_code,400,response.data)
        self.assertEqual(str(response.data['data']['row']),'3')
        self.assertFalse(Post.objects.filter(code__in=['valid','null']).exists())

    def test_reserved_literal_string_primary_roundtrips(self):
        area=Area.objects.create(code='#NULL',name='原名称')
        book=self.template('area',update=True)
        sheet=book['数据'];columns={c.value:c.column for c in sheet[1]}
        row=next(r for r in range(2,sheet.max_row+1) if sheet.cell(r,columns[UPDATE_KEY]).value=='#TEXT:#NULL')
        sheet.cell(row,columns['名称'],'新名称')
        response=self.upload('area',book)
        self.assertEqual(response.status_code,200,response.data)
        area.refresh_from_db();self.assertEqual(area.name,'新名称')

    def test_synchronous_and_task_export_share_results_and_digest(self):
        Post.objects.create(name='异步岗位', code='001', sort=0)
        sync = self.export_book('post', code='001')
        with patch('application.celery.app.send_task') as submit:
            submit.return_value.id = 'excel-task'
            response = self.client.post('/api/system/post/async_export/?code=001', {}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        args = submit.call_args.kwargs['args']
        with override_settings(MEDIA_ROOT=Path(__file__).resolve().parents[2]/'media'/'test-excel'):
            export_file.apply(args=args, throw=True)
            item = DownloadCenter.objects.get(pk=args[0])
            content = Path(item.url.path).read_bytes()
            self.assertEqual((item.task_status, item.size, item.md5sum), (2, len(content), md5(content).hexdigest()))
            async_book = load_workbook(BytesIO(content))
            self.assertEqual(list(sync['数据'].values), list(async_book['数据'].values))
            async_book.close()
            item.url.delete(save=False)
        sync.close()
