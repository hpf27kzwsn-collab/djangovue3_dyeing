import json
from importlib import import_module
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from django.apps import apps
from django.core.management import call_command
from django.db import connection
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient

from system.excel_fields import UPDATE_KEY
from system.models import Users, Role, Menu, MenuField, FieldPermission, LoginLog


class TenthImplementationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('init_data', verbosity=0, stdout=StringIO())

    def setUp(self):
        self.admin = Users.objects.get(username='superadmin')
        self.client = APIClient()
        self.client.force_authenticate(self.admin)

    def login(self, username, password):
        client = APIClient()
        response = client.post('/api/login/', {'username': username, 'password': password}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        client.credentials(HTTP_AUTHORIZATION='Bearer ' + response.data['data']['access'])
        self.assertEqual(client.get('/api/account/').data['data']['user']['username'], username)
        return response.data['data']

    def workbook(self, update=False, **filters):
        response = self.client.get('/api/system/user/' + ('update_template/' if update else 'import_template/'), filters)
        self.assertEqual(response.status_code, 200)
        content = b''.join(response.streaming_content)
        return load_workbook(BytesIO(content))

    def upload(self, book):
        output = BytesIO()
        book.save(output)
        book.close()
        response = self.client.post('/api/system/user/import/', {'file': SimpleUploadedFile('用户.xlsx', output.getvalue())}, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)

    def test_plain_username_create_update_login_unique_and_model_length(self):
        response = self.client.post('/api/system/user/', {'username': '部门 用户:一', 'password': 'initial'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        pk = response.data['data']['id']
        self.login('部门 用户:一', 'initial')
        renamed = '改名 用户:二'
        response = self.client.patch(f'/api/system/user/{pk}/', {'username': renamed}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.login(renamed, 'initial')
        duplicate = self.client.post('/api/system/user/', {'username': renamed}, format='json')
        self.assertEqual(duplicate.status_code, 400)
        long_name = '用' * 149 + ':'
        response = self.client.post('/api/system/user/', {'username': long_name, 'name': '长账号'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        response = self.client.patch(f'/api/system/user/{pk}/', {'username': '用' * 151}, format='json')
        self.assertEqual(response.status_code, 400)

    def test_plain_username_xlsx_import_update_and_login(self):
        book = self.workbook()
        for cell in book['数据'][1]:
            if cell.value == '登录账号':
                book['数据'].cell(2, cell.column, 'Excel 用户:一')
            elif cell.value == '用户名称':
                book['数据'].cell(2, cell.column, 'Excel导入')
        self.upload(book)
        user = Users.objects.get(username='Excel 用户:一')
        self.login(user.username, 'admin123456')
        book = self.workbook(update=True, id=user.pk)
        columns = {cell.value: cell.column for cell in book['数据'][1]}
        book['数据'].cell(2, columns[UPDATE_KEY], user.pk)
        book['数据'].cell(2, columns['登录账号'], 'Excel 用户:二')
        self.upload(book)
        user.refresh_from_db()
        self.assertEqual(user.username, 'Excel 用户:二')
        self.login(user.username, 'admin123456')

    def test_plain_username_business_initializer_and_repeat_run(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'users.json'
            path.write_text(json.dumps([{'username': '初始化 用户:一', 'name': '初始化'}]), encoding='utf-8')
            for _ in range(2):
                call_command('load_init_data', 'system.Users', str(path), unique_fields=['username'], stdout=StringIO())
        self.assertEqual(Users.objects.filter(username='初始化 用户:一').count(), 1)
        self.login('初始化 用户:一', 'admin123456')

    def test_both_admin_resets_preserve_first_password_count(self):
        for initial_count in (0, 4):
            user = Users.objects.create_user(username=f'重置 用户:{initial_count}', password='before', pwd_change_count=initial_count)
            response = self.client.post(f'/api/system/user/{user.pk}/reset_password/', {'password': 'changed', 'password_confirm': 'changed'}, format='json')
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(self.login(user.username, 'changed')['pwd_change_count'], initial_count)
            response = self.client.put(f'/api/system/user/{user.pk}/reset_to_default_password/', {}, format='json')
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(self.login(user.username, 'admin123456')['pwd_change_count'], initial_count)
            user.refresh_from_db()
            self.assertEqual(user.pwd_change_count, initial_count)

    def test_unused_role_field_cleanup_preserves_other_grants_and_initializer(self):
        role = Role.objects.get(key='admin')
        menu = Menu.objects.get(component_name='role')
        obsolete = MenuField.objects.create(menu=menu, model='system.Role', field_name='admin', title='管理员角色')
        FieldPermission.objects.create(role=role, field=obsolete)
        retained = list(FieldPermission.objects.exclude(field=obsolete).order_by('pk').values())
        migration = import_module('system.migrations.0019_windows_parity_cleanup')
        migration.remove_unused_role_field(apps, SimpleNamespace(connection=connection))
        self.assertFalse(MenuField.objects.filter(pk=obsolete.pk).exists())
        self.assertEqual(list(FieldPermission.objects.order_by('pk').values()), retained)
        call_command('init_data', verbosity=0, stdout=StringIO())
        fields = self.client.get('/api/meta/role/').data['data']['fields']
        self.assertNotIn('admin', [field['key'] for field in fields])
        self.assertEqual(LoginLog._meta.get_field('login_type').choices, [(1, '账号密码')])
