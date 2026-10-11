from io import BytesIO
from unittest.mock import patch

from django.core.management import call_command
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from system.models import Users, Post, Role, SystemConfig, Dictionary, Dept, MenuButton, RoleMenuButtonPermission


class FifthAuditTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('init_data', verbosity=0)

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(Users.objects.get(username='superadmin'))

    def test_exclusion_and_virtual_fields_in_api_and_xlsx(self):
        post = Post.objects.create(name='岗位', code='ERP')
        options = {'serializer_exclude_fields':['code'], 'fields':{
            'code':{'title':'编码'}, 'alias':{'source':'name', 'title':'岗位别名'}}}
        with patch('core.api.resource_options', return_value=options), patch('system.excel_fields.resource_options', return_value=options):
            data = self.client.get(f'/api/system/post/{post.pk}/').data['data']
            self.assertNotIn('code', data)
            self.assertEqual(data['alias'], '岗位')
            response = self.client.get('/api/system/post/export/')
            self.assertEqual(response.status_code, 200)
            book = load_workbook(BytesIO(b''.join(response.streaming_content)))
            values = list(book['数据'].values)
            self.assertNotIn('编码', values[0])
            self.assertEqual(values[1][values[0].index('岗位别名')], '岗位')
            book.close()

    def test_user_export_omits_excluded_optional_column(self):
        with patch('core.api.resource_options', return_value={'serializer_exclude_fields':['mobile']}):
            response = self.client.get('/api/system/user/export/')
            self.assertEqual(response.status_code, 200)
            book = load_workbook(BytesIO(b''.join(response.streaming_content)))
            self.assertNotIn('手机号码', [cell.value for cell in book['数据'][1]])
            book.close()

    def test_related_boolean_filters_are_typed_and_distinct(self):
        user = Users.objects.create_user(username='related-bool')
        user.role.add(Role.objects.create(name='启用一', key='active-one'), Role.objects.create(name='启用二', key='active-two'))
        for key in ['role__status', 'role__status__in']:
            response = self.client.get('/api/system/user/', {key:'true', 'username':user.username})
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(response.data['total'], 1)

    def test_configured_business_key_reverse_lookup(self):
        group = SystemConfig.objects.create(key='reverse-test', title='反向关联')
        SystemConfig.objects.create(parent=group, key='chosen', title='岗位', form_item_type=13,
            value='ERP-CODE', setting={'table':'post', 'primarykey':'code'})
        response = self.client.get('/api/system/system_config/relation_info/',
            {'var_name':'chosen', 'table':'post', 'relation_ids':'ERP-CODE'})
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['data'][0]['group']['id'], group.pk)

    def test_association_list_has_target_action_context(self):
        group = SystemConfig.objects.create(key='action-test', title='列表上下文')
        Dictionary.objects.create(label='字典组', value='action-test')
        item = SystemConfig.objects.create(parent=group, key='chosen', title='字典', form_item_type=13,
            setting={'table':'dictionary', 'primarykey':'id'})
        response = self.client.get(f'/api/system/system_config/{item.pk}/table_data/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(all(row['parent'] is None for row in response.data['data']))

    def test_public_prefix_does_not_expand_empty_segment(self):
        group = SystemConfig.objects.create(key='private-prefix', title='其他公开项')
        SystemConfig.objects.create(parent=group, key='value', title='值', value='visible', status=True)
        self.client.force_authenticate(None)
        response = self.client.get('/api/config/values/', {'key':'site.|'})
        self.assertEqual(response.status_code, 200, response.data)
        self.assertNotIn('private-prefix.value', response.data['data'])

    def test_public_brand_ignores_private_config_but_login_policy_remains(self):
        title = SystemConfig.objects.get(parent__key='site', key='title')
        title.value, title.status = 'PRIVATE-TITLE', False
        title.save()
        self.client.force_authenticate(None)
        data = self.client.get('/api/config/').data['data']
        self.assertEqual(data['title'], 'ERP 管理平台')
        self.assertTrue(data['first_password'])
        self.assertIn('copyright', data)

    def test_anonymous_dictionary_read_does_not_grant_crud(self):
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get('/api/dictionary/', {'key':'status'}).status_code, 200)
        self.assertEqual(self.client.post('/api/system/dictionary/', {'label':'禁止写入','value':'write'}, format='json').status_code, 401)

    def test_department_dependency_works_without_department_list_permission(self):
        department = Dept.objects.create(name='可查询部门')
        other = Dept.objects.create(name='其他部门')
        role = Role.objects.create(name='只读用户', key='user-only')
        user = Users.objects.create_user(username='user-only', dept=department, dept_belong=department)
        user.role.add(role)
        button = MenuButton.objects.get(value='user:list')
        RoleMenuButtonPermission.objects.create(role=role, menu_button=button, data_range=2)
        Users.objects.create_user(username='other-dept', dept=other, dept_belong=other)
        self.client.force_authenticate(user)
        self.assertEqual(self.client.get('/api/system/dept/').status_code, 403)
        self.assertEqual(self.client.get('/api/system/dept/options/').status_code, 200)
        response = self.client.get(f'/api/system/dept/{department.pk}/info/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['data']['count'], 1)
        response = self.client.get(f'/api/system/dept/{other.pk}/info/')
        self.assertEqual(response.data['data']['count'], 0)

    def test_member_search_parameters_do_not_filter_the_parent_role(self):
        role = Role.objects.create(name='父角色', key='member-search')
        member = Users.objects.create_user(username='search-member', name='成员姓名')
        member.role.add(role)
        for name in ['', '成员姓名']:
            response = self.client.get(f'/api/system/role/{role.pk}/users/',
                {'authorized':'1','name':name,'page':1,'limit':20})
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(response.data['data'][0]['id'], member.pk)
