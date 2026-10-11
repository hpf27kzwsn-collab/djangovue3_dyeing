from io import BytesIO
from unittest import skipUnless
from unittest.mock import patch

from django.apps import apps
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from core.api import serializer_for
from core.permissions import field_access
from system.excel_fields import supports_import
from system.models import Users, Role, MenuField, FieldPermission, MenuButton, RoleMenuButtonPermission, DownloadCenter
from system.tasks import export_file


@skipUnless(apps.is_installed('system.tests.twin_a'), 'requires extension_settings')
class BuiltinNameTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('init_data', verbosity=0)

    def setUp(self):
        self.client = APIClient()
        self.admin = Users.objects.get(username='superadmin')
        self.client.force_authenticate(self.admin)

    def workbook(self, entity, action, **query):
        response = self.client.get(f'/api/system/{entity}/{action}/', query)
        self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
        book = load_workbook(BytesIO(b''.join(response.streaming_content)))
        return book

    def upload(self, entity, book):
        content = BytesIO()
        book.save(content)
        book.close()
        return self.client.post(f'/api/system/{entity}/import/',
            {'file': SimpleUploadedFile('同名业务.xlsx', content.getvalue())}, format='multipart')

    def test_same_builtin_names_metadata_crud_and_xlsx_roundtrip(self):
        from system.tests.twin_a import models
        parent = models.Twin.objects.create(name='父项')
        tag = models.Twin.objects.create(name='标签')
        for name in ['Users', 'LoginLog', 'OperationLog', 'DownloadCenter', 'TaskResult', 'PlainEntry']:
            with self.subTest(model=name):
                model = getattr(models, name)
                entity = f'identity_a_{model._meta.model_name}'
                url = f'/api/system/{entity}/'
                meta = self.client.get(f'/api/meta/{entity}/')
                self.assertEqual(meta.status_code, 200, meta.data)
                self.assertTrue(meta.data['data']['excel_import'])
                created = self.client.post(url, {'name': name, 'quantity': 7, 'dept_owner': '业务负责人',
                    'parent': parent.pk, 'tags': [tag.pk]}, format='json')
                self.assertEqual(created.status_code, 200, created.data)
                item = model.objects.get(pk=created.data['data']['id'])
                self.assertEqual(self.client.get(url).status_code, 200)
                template = self.workbook(entity, 'import_template')
                titles = [cell.value for cell in template['数据'][1]]
                for title in ['业务名称', '数量', '业务负责人', '业务父项', '业务标签']:
                    self.assertIn(title, titles)
                record = {'业务名称': '导入'+name, '数量': 0, '业务负责人': '导入负责人',
                    '业务父项': f'父项 [{parent.pk}]', '业务标签': f'标签 [{tag.pk}]'}
                for cell in template['数据'][1]:
                    template['数据'].cell(2, cell.column).value = record.get(cell.value)
                imported = self.upload(entity, template)
                self.assertEqual(imported.status_code, 200, imported.data)
                second = model.objects.get(name='导入'+name)
                self.assertEqual((second.quantity, second.parent_id, second.dept_owner), (0, parent.pk, '导入负责人'))
                self.assertEqual(list(second.tags.values_list('pk', flat=True)), [tag.pk])
                update = self.workbook(entity, 'update_template', id=item.pk)
                quantity_column = next(cell.column for cell in update['数据'][1] if cell.value == '数量')
                update['数据'].cell(2, quantity_column).value = 11
                updated = self.upload(entity, update)
                self.assertEqual(updated.status_code, 200, updated.data)
                item.refresh_from_db()
                self.assertEqual(item.quantity, 11)
                sync = self.workbook(entity, 'export')
                rows = list(sync['数据'].values)
                self.assertIn('数量', rows[0]);self.assertIn('业务负责人', rows[0])
                sync.close()
                deleted = self.client.delete(url+str(second.pk)+'/')
                self.assertEqual(deleted.status_code, 200)

    def test_same_name_field_permissions_are_separate_from_internal_users(self):
        from system.tests.twin_a.models import Users as BusinessUsers
        role = Role.objects.create(name='业务用户角色', key='builtin-collision')
        actor = Users.objects.create_user(username='builtin-actor')
        actor.role.add(role)
        for action in ['list', 'export-get', 'import-template-get']:
            RoleMenuButtonPermission.objects.create(role=role,
                menu_button=MenuButton.objects.get(value=f'identity_a_users:{action}'), data_range=3)
        field = MenuField.objects.get(model='identity_a.Users', field_name='quantity')
        FieldPermission.objects.create(role=role, field=field, is_query=True, is_create=False, is_update=False)
        self.assertTrue(field_access(actor, BusinessUsers)['quantity']['is_query'])
        BusinessUsers.objects.create(name='普通角色可见数量', quantity=23)
        self.client.force_authenticate(actor)
        result = self.client.get('/api/system/identity_a_users/')
        self.assertEqual(result.status_code, 200, result.data)
        self.assertIn('quantity', result.data['data'][0])
        self.assertNotIn('name', result.data['data'][0])
        book = self.workbook('identity_a_users', 'export')
        self.assertIn('数量', list(book['数据'].values)[0])
        self.assertNotIn('业务名称', list(book['数据'].values)[0])
        book.close()
        internal = serializer_for(Users)(context={'request': type('Request', (), {'user': self.admin})()})
        self.assertNotIn('password', internal.fields)
        self.assertNotIn('groups', internal.fields)
        for entity in ['login_log', 'operation_log', 'download_center', 'task_result']:
            from system.catalog import ENTITIES
            self.assertFalse(supports_import(ENTITIES[entity]))

    def test_same_name_async_export_uses_full_business_columns(self):
        from system.tests.twin_a.models import Users as BusinessUsers
        BusinessUsers.objects.create(name='异步业务用户', quantity=31, dept_owner='业务负责人')
        with patch('application.celery.app.send_task') as send:
            send.return_value.id = 'builtin-name-test-task'
            result = self.client.post('/api/system/identity_a_users/async_export/', {}, format='json')
        self.assertEqual(result.status_code, 200, result.data)
        args = send.call_args.kwargs['args']
        from tempfile import TemporaryDirectory
        with TemporaryDirectory() as media, override_settings(MEDIA_ROOT=media):
            export_file.run(*args)
            item = DownloadCenter.objects.get(pk=result.data['data']['id'])
            self.assertEqual(item.task_status, 2)
            with item.url.open('rb') as stream:
                book = load_workbook(stream)
                rows = list(book['数据'].values)
                self.assertEqual(rows[1][rows[0].index('数量')], 31)
                self.assertEqual(rows[1][rows[0].index('业务负责人')], '业务负责人')
                book.close()
