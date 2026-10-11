from io import BytesIO
from django.core.management import call_command
from django.test import TestCase
from rest_framework.test import APIClient
from openpyxl import load_workbook
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from pathlib import Path
from system.catalog import ENTITIES
from system.models import Users, Dept, Role, Post, Menu, MenuButton, MenuField, RoleMenuButtonPermission, FieldPermission


class PlatformTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('init_data', verbosity=0)

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(Users.objects.get(username='superadmin'))

    def test_login_account_and_metadata(self):
        self.client.force_authenticate(None)
        response = self.client.post('/api/login/', {'username':'superadmin','password':'admin123456'}, format='json')
        self.assertEqual(response.status_code, 200)
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+response.data['data']['access'])
        self.assertTrue(self.client.get('/api/account/').data['data']['superuser'])
        self.assertEqual(self.client.get('/api/meta/user/').data['data']['title'], '用户')

    def test_create_user_hashes_password_and_patch(self):
        response = self.client.post('/api/system/user/', {'username':'new-user','password':'hello','name':'小李'}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        user = Users.objects.get(username='new-user')
        self.assertTrue(user.check_password('hello'))
        self.assertNotIn('password', response.data['data'])
        response = self.client.patch(f'/api/system/user/{user.pk}/', {'name':'小王'}, format='json')
        self.assertEqual(response.data['data']['name'], '小王')

    def test_data_scope_applies_to_bulk_and_export(self):
        department = Dept.objects.create(name='研发')
        other = Dept.objects.create(name='采购')
        user = Users.objects.create_user(username='staff', password='hello', dept=department)
        role = Role.objects.create(name='研发角色', key='staff')
        user.role.add(role)
        own = Post.objects.create(name='工程师', code='engineering', creator=user, dept_belong=department)
        foreign = Post.objects.create(name='采购员', code='purchasing', dept_belong=other)
        menu = Menu.objects.get(component_name='post')
        for key, verb, route in [('list',0,''), ('export',0,'export/'),('bulk',3,'multiple_delete/')]:
            button = MenuButton.objects.create(name=key, value='staff:'+key, menu=menu, method=verb, api='/api/system/post/'+route)
            RoleMenuButtonPermission.objects.create(role=role, menu_button=button, data_range=2)
        for field in MenuField.objects.filter(model='system.Post'):
            FieldPermission.objects.create(role=role, field=field, is_query=field.field_name!='code')
        self.client.force_authenticate(user)
        response = self.client.get('/api/system/post/')
        self.assertEqual([row['id'] for row in response.data['data']], [own.pk])
        self.assertNotIn('code', response.data['data'][0])
        exported = self.client.get('/api/system/post/export/')
        sheet = load_workbook(BytesIO(b''.join(exported.streaming_content))).active
        self.assertNotIn('岗位编码', [cell.value for cell in sheet[1]])
        self.client.delete('/api/system/post/multiple_delete/', {'keys':[own.pk, foreign.pk]}, format='json')
        self.assertFalse(Post.objects.filter(pk=own.pk).exists())
        self.assertTrue(Post.objects.filter(pk=foreign.pk).exists())

    def test_role_permissions_replace_and_message_delivery(self):
        role = Role.objects.get(key='admin')
        response = self.client.get(f'/api/system/role/{role.pk}/permissions/')
        self.assertEqual(response.status_code, 200, response.data)
        user = Users.objects.get(username='superadmin')
        message = self.client.post('/api/system/message_center/', {'title':'通知','content':'欢迎','target_user':[user.pk]}, format='json')
        self.assertEqual(message.status_code, 200, message.data)
        published = self.client.post(f"/api/system/message_center/{message.data['data']['id']}/publish/", {}, format='json')
        self.assertEqual(published.data['data']['count'], 1)
        inbox = self.client.get('/api/inbox/').data['data']
        self.assertEqual(inbox['unread'], 1)
        self.client.post('/api/inbox/', {'ids':[inbox['messages'][0]['id']]}, format='json')
        self.assertEqual(self.client.get('/api/inbox/').data['data']['unread'], 0)

    def test_denied_field_write(self):
        user = Users.objects.create_user(username='writer', password='hello')
        role = Role.objects.create(name='只读编码', key='writer')
        user.role.add(role)
        button = MenuButton.objects.get(value='post:add')
        RoleMenuButtonPermission.objects.create(role=role, menu_button=button, data_range=3)
        for field in MenuField.objects.filter(model='system.Post'):
            FieldPermission.objects.create(role=role, field=field, is_create=field.field_name!='code')
        self.client.force_authenticate(user)
        response = self.client.post('/api/system/post/', {'name':'工程师','code':'restricted'}, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Post.objects.filter(code='restricted').exists())

    def test_all_metadata_and_crontab_timezone(self):
        for entity in ENTITIES:
            with self.subTest(entity=entity):
                self.assertEqual(self.client.get('/api/meta/'+entity+'/').status_code,200)
        response=self.client.post('/api/system/crontab_schedule/',{'timezone':'Asia/Shanghai'},format='json')
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.data['data']['timezone'],'Asia/Shanghai')
        response=self.client.get('/api/system/crontab_schedule/')
        self.assertEqual(response.data['data'][0]['timezone'],'Asia/Shanghai')

    def test_sse_accepts_event_stream_content_negotiation(self):
        response=self.client.get('/api/events/',HTTP_ACCEPT='text/event-stream')
        self.assertEqual(response.status_code,200)
        self.assertTrue(response.streaming)
        self.assertEqual(response['Content-Type'],'text/event-stream')
        # 未消费的 SSE 响应按 Django 测试客户端的关闭流程清理，保留外层测试事务。
        from django.test.client import closing_iterator_wrapper
        list(closing_iterator_wrapper((),response.close))

    def test_schedule_requires_a_rule(self):
        response=self.client.post('/api/system/periodic_task/',{'name':'无规则','task':'system.tasks.demo_task'},format='json')
        self.assertEqual(response.status_code,400)

    def test_account_is_available_without_user_management_permission(self):
        user=Users.objects.create_user(username='operator',password='hello',name='操作员')
        self.client.force_authenticate(user)
        account=self.client.get('/api/account/')
        self.assertEqual(account.data['data']['user']['username'],'operator')
        response=self.client.patch('/api/account/',{'name':'新的姓名'},format='json')
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.data['data']['user']['name'],'新的姓名')
        self.assertEqual(self.client.get('/api/system/user/').status_code,403)

    def test_role_grants_can_be_saved(self):
        from system.models import RoleMenuPermission
        role=Role.objects.create(name='测试角色',key='role-save')
        menu=Menu.objects.get(component_name='post')
        button=MenuButton.objects.get(value='post:list')
        column=MenuField.objects.get(model='system.Post',field_name='code')
        payload={'menus':[menu.pk],'buttons':[{'id':button.pk,'data_range':2,'dept':[]}],
                 'fields':[{'id':column.pk,'is_query':True,'is_create':False,'is_update':False}]}
        response=self.client.put(f'/api/system/role/{role.pk}/permissions/',payload,format='json')
        self.assertEqual(response.status_code,200,response.data)
        self.assertTrue(RoleMenuPermission.objects.filter(role=role,menu=menu).exists())
        self.assertEqual(RoleMenuButtonPermission.objects.get(role=role).data_range,2)
        self.assertFalse(FieldPermission.objects.get(role=role).is_update)

    def test_metadata_defaults_form_payload(self):
        meta=self.client.get('/api/meta/post/').data['data']
        payload={f['key']:f['default'] for f in meta['fields'] if not f['readonly']}
        payload.update(name='表单默认值',code='form-default')
        response=self.client.post('/api/system/post/',payload,format='json')
        self.assertEqual(response.status_code,200,response.data)

    def test_initialization_preserves_edited_admin_permissions(self):
        role=Role.objects.get(key='admin')
        grant=RoleMenuButtonPermission.objects.get(role=role,menu_button__value='post:add')
        button=grant.menu_button
        grant.delete()
        call_command('init_data',verbosity=0)
        self.assertFalse(RoleMenuButtonPermission.objects.filter(role=role,menu_button=button).exists())

    def test_restql_select_and_alias_respects_fields(self):
        user=Users.objects.create_user(username='query-user',password='hello')
        role=Role.objects.create(name='动态查询角色',key='query')
        user.role.add(role)
        Post.objects.create(name='研发',code='secret')
        RoleMenuButtonPermission.objects.create(role=role,menu_button=MenuButton.objects.get(value='post:list'),data_range=3)
        for field in MenuField.objects.filter(model='system.Post'):
            FieldPermission.objects.create(role=role,field=field,is_query=field.field_name!='code')
        self.client.force_authenticate(user)
        response=self.client.get('/api/system/post/',{'query':'{id, label:name, exposed:code}'})
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(response.data['data'][0]['label'],'研发')
        self.assertNotIn('exposed',response.data['data'][0])

    def test_excel_roundtrip(self):
        response=self.client.get('/api/system/post/import_template/')
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(b''.join(response.streaming_content)))
        sheet=book['数据']
        values={'岗位名称':'设计师','岗位编码':'designer'}
        for index,cell in enumerate(sheet[1],1):
            if cell.value in values:
                sheet.cell(2,index,values[cell.value])
        output=BytesIO();book.save(output);book.close()
        upload=SimpleUploadedFile('posts.xlsx',output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        imported=self.client.post('/api/system/post/import/',{'file':upload},format='multipart')
        self.assertEqual(imported.status_code,200,imported.data)
        self.assertEqual(Post.objects.filter(code='designer').count(),1)

    @override_settings(MEDIA_ROOT=Path(__file__).resolve().parents[2]/'media'/'test-platform')
    def test_file_upload_and_delete(self):
        upload=SimpleUploadedFile('hello.txt',b'hello',content_type='text/plain')
        response=self.client.post('/api/system/file/',{'url':upload},format='multipart')
        self.assertEqual(response.status_code,200,response.data)
        from system.models import FileList
        item=FileList.objects.get(pk=response.data['data']['id'])
        path=Path(item.url.path)
        self.assertTrue(path.exists())
        self.assertEqual(item.size,5)
        self.assertEqual(item.file_type,3)
        updated=self.client.patch(f'/api/system/file/{item.pk}/',{'url':SimpleUploadedFile('new.txt',b'replaced',content_type='text/plain')},format='multipart')
        self.assertEqual(updated.status_code,200,updated.data)
        self.assertFalse(path.exists())
        item.refresh_from_db()
        path=Path(item.url.path)
        self.assertEqual(item.size,8)
        self.client.delete(f'/api/system/file/{item.pk}/')
        self.assertFalse(path.exists())

    def test_five_data_ranges_and_role_union(self):
        root=Dept.objects.create(name='研发')
        child=Dept.objects.create(name='研发一组',parent=root)
        foreign=Dept.objects.create(name='采购')
        user=Users.objects.create_user(username='scope',password='hello',dept=root)
        role=Role.objects.create(name='范围测试',key='scope')
        user.role.add(role)
        own=Post.objects.create(name='本人',code='own',creator=user,dept_belong=root)
        colleague=Post.objects.create(name='同部门',code='colleague',dept_belong=root)
        descendant=Post.objects.create(name='下级',code='child',dept_belong=child)
        outside=Post.objects.create(name='其他部门',code='foreign',dept_belong=foreign)
        grant=RoleMenuButtonPermission.objects.create(role=role,menu_button=MenuButton.objects.get(value='post:list'))
        grant.dept.add(foreign)
        self.client.force_authenticate(user)
        for scope,expected in [(0,{own.pk}),(1,{own.pk,colleague.pk,descendant.pk}),(2,{own.pk,colleague.pk}),(3,{own.pk,colleague.pk,descendant.pk,outside.pk}),(4,{outside.pk})]:
            with self.subTest(scope=scope):
                grant.data_range=scope;grant.save()
                self.assertEqual({x['id'] for x in self.client.get('/api/system/post/').data['data']},expected)
        second=Role.objects.create(name='附加角色',key='scope2')
        user.role.add(second)
        RoleMenuButtonPermission.objects.create(role=second,menu_button=grant.menu_button,data_range=0)
        self.assertEqual({x['id'] for x in self.client.get('/api/system/post/').data['data']},{own.pk,outside.pk})

    def test_captcha_refresh_password_and_unlock(self):
        from captcha.models import CaptchaStore
        from system.models import SystemConfig
        SystemConfig.objects.filter(parent__key='login',key='captcha').update(value=True)
        self.client.force_authenticate(None)
        self.client.credentials(HTTP_AUTHORIZATION='Bearer expired-token')
        challenge=self.client.get('/api/captcha/').data['data']
        answer=CaptchaStore.objects.get(hashkey=challenge['key']).response
        response=self.client.post('/api/login/',{'username':'superadmin','password':'admin123456','captcha_key':challenge['key'],'captcha':answer},format='json')
        self.assertEqual(response.status_code,200,response.data)
        tokens=response.data['data']
        self.client.credentials()
        refresh=self.client.post('/api/token/refresh/',{'refresh':tokens['refresh']},format='json')
        self.assertEqual(refresh.status_code,200,refresh.data)
        self.client.credentials(HTTP_AUTHORIZATION='Bearer '+refresh.data['access'])
        self.assertEqual(self.client.post('/api/account/',{'password':'admin123456'},format='json').status_code,200)
        changed=self.client.post('/api/password/',{'old_password':'admin123456','new_password':'new-password','new_password_confirm':'new-password'},format='json')
        self.assertEqual(changed.status_code,200,changed.data)
        user=Users.objects.get(username='superadmin')
        self.assertTrue(user.check_password('new-password'))
        self.assertEqual(user.pwd_change_count,1)

    def test_associated_table_uses_related_data_scope(self):
        from system.models import SystemConfig
        own=Dept.objects.create(name='本部门')
        other=Dept.objects.create(name='其他部门')
        user=Users.objects.create_user(username='relation',password='hello',dept=own)
        role=Role.objects.create(name='关联数据角色',key='relation');user.role.add(role)
        post=Post.objects.create(name='允许关联',code='allowed',dept_belong=own)
        Post.objects.create(name='其他数据',code='other',dept_belong=other)
        config=SystemConfig.objects.create(title='关联岗位',key='relation.post',setting={'table':'post'},dept_belong=own)
        RoleMenuButtonPermission.objects.create(role=role,menu_button=MenuButton.objects.get(value='post:list'),data_range=2)
        RoleMenuButtonPermission.objects.create(role=role,menu_button=MenuButton.objects.get(value='system_config:table-data-get'),data_range=2)
        self.client.force_authenticate(user)
        response=self.client.get(f'/api/system/system_config/{config.pk}/table_data/')
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual([x['id'] for x in response.data['data']],[post.pk])

    def test_password_flows_and_audit_does_not_store_passwords(self):
        from system.models import SystemConfig, OperationLog
        user=Users.objects.create_user(username='initial-password',password='old')
        SystemConfig.objects.filter(parent__key='base',key='default_password').update(value='configured-password')
        response=self.client.put(f'/api/system/user/{user.pk}/reset_to_default_password/',{},format='json')
        self.assertEqual(response.status_code,200,response.data)
        user.refresh_from_db();self.assertTrue(user.check_password('configured-password'))
        self.client.force_authenticate(user)
        response=self.client.post('/api/password/first/',{'new_password':'mine','new_password_confirm':'different'},format='json')
        self.assertEqual(response.status_code,400)
        response=self.client.post('/api/password/first/',{'new_password':'mine','new_password_confirm':'mine'},format='json')
        self.assertEqual(response.status_code,200,response.data)
        user.refresh_from_db();self.assertTrue(user.check_password('mine'));self.assertEqual(user.pwd_change_count,1)
        log=OperationLog.objects.filter(request_path='/api/password/first/').latest('pk')
        self.assertEqual(log.request_body,{})
        self.assertEqual(log.response_code,2000)
        self.assertEqual(log.json_result['code'],2000)
        self.assertEqual(log.request_modular,'账户与系统')

    def test_field_matching_and_button_generation_are_repeatable(self):
        menu=Menu.objects.create(name='测试菜单',component_name='auditdemo',web_path='/system/auditdemo')
        payload={'menu':menu.pk,'model':'system.Post'}
        for _ in range(2):
            response=self.client.post('/api/system/column/auto_match_fields/',payload,format='json')
            self.assertEqual(response.status_code,200,response.data)
            response=self.client.post('/api/system/menu_button/batch_create/',{'menu':menu.pk,'entity':'post'},format='json')
            self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(MenuField.objects.filter(menu=menu).count(),len(Post._meta.fields))
        self.assertGreater(MenuButton.objects.filter(menu=menu).count(),9)
        self.assertEqual(MenuButton.objects.get(menu=menu,value='auditdemo:export-get').api,'/api/system/post/export/')
        self.assertTrue(MenuButton.objects.filter(menu=menu,value='auditdemo:batch-update-put').exists())

    def test_messages_deliver_on_creation_to_each_target_type(self):
        from system.models import MessageCenterTargetUser
        department=Dept.objects.create(name='消息部门')
        role=Role.objects.create(name='消息角色',key='message-test')
        receiver=Users.objects.create_user(username='message-user',password='hello',dept=department)
        receiver.role.add(role)
        for target,payload in [(0,{'target_user':[receiver.pk]}),(1,{'target_role':[role.pk]}),(2,{'target_dept':[department.pk]}),(3,{})]:
            response=self.client.post('/api/system/message_center/',{'title':str(target),'content':'通知','target_type':target,**payload},format='json')
            self.assertEqual(response.status_code,200,response.data)
            message_id=response.data['data']['id']
            self.assertTrue(MessageCenterTargetUser.objects.filter(messagecenter_id=message_id,users=receiver).exists())
            self.assertEqual(MessageCenterTargetUser.objects.filter(messagecenter_id=message_id).count(),Users.objects.count() if target==3 else 1)
        self.client.force_authenticate(receiver)
        response=self.client.get('/api/system/message_center/get_self_receive/')
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(len(response.data['data']),4)
        self.assertEqual(self.client.get('/api/system/message_center/get_newest_msg/').data['data']['title'],'3')

    def test_downloads_remain_owned_and_move_handles_equal_sorts(self):
        from system.models import DownloadCenter
        user=Users.objects.create_user(username='download-user',password='hello')
        role=Role.objects.create(name='下载角色',key='download-test');user.role.add(role)
        RoleMenuButtonPermission.objects.create(role=role,menu_button=MenuButton.objects.get(value='download_center:list'),data_range=3)
        own=DownloadCenter.objects.create(task_name='自己',creator=user)
        DownloadCenter.objects.create(task_name='别人')
        self.client.force_authenticate(user)
        self.assertEqual([x['id'] for x in self.client.get('/api/system/download_center/').data['data']],[own.pk])
        self.client.force_authenticate(Users.objects.get(username='superadmin'))
        parent=Dept.objects.create(name='排序根')
        first=Dept.objects.create(name='一',parent=parent);second=Dept.objects.create(name='二',parent=parent)
        response=self.client.post('/api/system/dept/move_up/',{'id':second.pk},format='json')
        self.assertEqual(response.status_code,200,response.data)
        self.assertEqual(list(Dept.objects.filter(parent=parent).order_by('sort','pk').values_list('pk',flat=True)),[first.pk,second.pk])

    def test_area_derives_fields_and_dictionary_types(self):
        parent=self.client.post('/api/system/area/',{'code':'test-province','name':'安徽省'},format='json')
        self.assertEqual(parent.status_code,200,parent.data)
        child=self.client.post('/api/system/area/',{'code':'test-city','name':'合肥市','pcode':'test-province'},format='json')
        self.assertEqual(child.status_code,200,child.data)
        self.assertEqual(child.data['data']['pinyin'],'hefeishi')
        self.assertEqual(child.data['data']['initials'],'H')
        self.assertEqual(child.data['data']['level'],2)
        dictionary=self.client.get('/api/meta/dictionary/').data['data']
        self.assertEqual(len(next(x for x in dictionary['fields'] if x['key']=='type')['options']),8)
        self.assertIs(self.client.get('/api/dictionary/',{'key':'status'}).data['data'][0]['value'],True)

    def test_group_configuration_visibility_and_uniqueness(self):
        from system.models import SystemConfig
        from system.auth import config
        base=SystemConfig.objects.get(key='base',parent=None)
        SystemConfig.objects.filter(parent=base,key='default_password').update(value='backend-only',status=False)
        self.assertEqual(config('base.default_password'),'backend-only')
        self.client.force_authenticate(None)
        self.client.credentials(HTTP_AUTHORIZATION='Bearer expired')
        response=self.client.get('/api/config/values/',{'key':'base|site'})
        self.assertEqual(response.status_code,200,response.data)
        self.assertNotIn('base.default_password',response.data['data'])
        self.assertIn('site.title',response.data['data'])
        self.client.credentials();self.client.force_authenticate(Users.objects.get(username='superadmin'))
        payload={'parent':base.pk,'key':'default_password','title':'重复'}
        response=self.client.post('/api/system/system_config/',payload,format='json')
        self.assertEqual(response.status_code,400,response.data)
        response=self.client.patch(f'/api/system/system_config/{base.pk}/',{'title':'基础'},format='json')
        self.assertEqual(response.status_code,200,response.data)

