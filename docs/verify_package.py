"""验证正在 8000 端口运行的桌面包服务及包内真实异步导出。"""
import json
import hashlib
import time
from io import BytesIO
from pathlib import Path
import requests
from openpyxl import load_workbook

base='http://127.0.0.1:8000'
session=requests.Session()
login=session.post(base+'/api/login/',json={'username':'superadmin','password':'admin123456'})
login.raise_for_status()
session.headers['Authorization']='Bearer '+login.json()['data']['access']
for path in ['/','/api/account/','/api/meta/periodic_task/','/api/system/area/','/api/schema/']:
    result=session.get(base+path);result.raise_for_status()
    print(path,result.status_code)
html=session.get(base+'/').text
import re
asset=re.search(r'src="([^"]+\.js)"',html).group(1)
assert session.get(base+asset).status_code==200
templates = {}
for mode in ['import_template', 'update_template']:
    response = session.get(base+f'/api/system/post/{mode}/', timeout=15)
    response.raise_for_status()
    book = load_workbook(BytesIO(response.content))
    headers = [cell.value for cell in book['数据'][1]]
    assert '岗位名称' in headers and '填写说明' in book.sheetnames and book['选项'].sheet_state == 'hidden'
    assert ('更新主键(勿改)' in headers) == (mode == 'update_template')
    templates[mode] = headers
    book.close()
sync = session.get(base+'/api/system/post/export/', timeout=15)
sync.raise_for_status()
sync_book = load_workbook(BytesIO(sync.content))
sync_rows = list(sync_book['数据'].values)
sync_book.close()
submitted=session.post(base+'/api/system/post/async_export/',json={})
submitted.raise_for_status()
item_id=submitted.json()['data']['id']
for _ in range(30):
    response=session.get(base+f'/api/system/download_center/{item_id}/')
    response.raise_for_status()
    item=response.json()['data']
    if item['task_status']==2:
        file=session.get(item['url']);file.raise_for_status()
        assert hashlib.md5(file.content).hexdigest()==item['md5sum']
        book=load_workbook(BytesIO(file.content))
        sheet=book['数据']
        assert '岗位名称' in [cell.value for cell in sheet[1]]
        assert list(sheet.values) == sync_rows
        assert len(file.content) == item['size']
        print('包内异步导出与同步数据一致',item['size'],'字节，表头：',[cell.value for cell in sheet[1]])
        book.close()
        break
    if item['task_status']==3:
        raise RuntimeError('包内异步任务执行失败')
    time.sleep(1)
else:
    raise RuntimeError('30 秒内包内导出没有完成')
account=session.get(base+'/api/account/').json()['data']
values=session.get(base+'/api/config/values/').json()['data']
assert 'base.default_password' not in values
with session.get(base+'/api/events/',headers={'Accept':'text/event-stream'},stream=True,timeout=15) as stream:
    stream.raise_for_status()
    assert stream.headers['Content-Type'].startswith('text/event-stream')
    event=next(line for line in stream.iter_lines(chunk_size=1) if line.startswith(b'data:'))
    assert 'unread' in json.loads(event[5:])
print('包内 SSE、公开配置与文件 MD5 验证通过')
Path('docs/package-verification.json').write_text(json.dumps({'login':200,'static_asset':asset,'download':item,
    'first_password':account['first_password'],'sse':'text/event-stream 200','download_md5_verified':True,
    'public_config_password_excluded':True,'xlsx_templates':templates,'sync_async_rows_equal':True,
    'download_size_verified':True},ensure_ascii=False,indent=2),encoding='utf-8')
