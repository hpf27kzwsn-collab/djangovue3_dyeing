"""独立库验证导航、报价xlsx原值承接及主从关系；不连接正式库。"""
import os
import sys
import json
import hashlib
from pathlib import Path
from collections import Counter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'application.settings')
import django
django.setup()
from django.apps import apps
from django.db import connection, transaction
from django.core.management import call_command
from django.test.utils import setup_test_environment
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient
from openpyxl import load_workbook
from system.models import Menu, MenuButton, MenuField, FieldPermission, RoleMenuButtonPermission, Users
from system.views import view_for
from plugins.erp_base.business_navigation import RESOURCES
from plugins.erp_base.foundation_catalog import DEFINITIONS
from plugins.erp_base.source_import import original_values
from plugins.erp_base.source_excel_history import columns
from plugins.erp_base.legacy_models import LegacyImportBatch
from plugins.erp_base.source_context import SourceDocumentContext
from plugins.erp_base.models import BusinessAction

assert str(connection.settings_dict['PORT']) == '5440' and connection.settings_dict['NAME'] == 'erp_master_xlsx_verify_20261021'
setup_test_environment()
actor = Users.objects.get(username='xlsx-update-20261021')
client = APIClient(HTTP_HOST='localhost');client.force_authenticate(actor)
work = ROOT / '.runtime-temp/sales-production-20261022'
pack = json.loads((work / 'quote-package.json').read_text(encoding='utf8'))
out = ROOT / 'outputs/染整报价导入包_20261022'
checks = []


def check(name, valid, detail=None):
    assert valid, (name, detail)
    checks.append({'name':name, 'passed':True, 'detail':detail})
    print('PASS', name, flush=True)


def digest(model):
    h=hashlib.sha256();count=0
    for row in model.objects.order_by('pk').values_list().iterator(chunk_size=1000):
        h.update(repr(tuple(bytes(v) if isinstance(v,memoryview) else v for v in row)).encode());count+=1
    return {'rows':count,'sha256':h.hexdigest()}


authorization_models=[MenuButton,MenuField,RoleMenuButtonPermission,FieldPermission,RoleMenuButtonPermission.dept.through]
before={m._meta.label:digest(m) for m in authorization_models}
data_before={label:digest(apps.get_model(label)) for label in RESOURCES}
menus_before={m.component_name:(m.pk,m.web_path,m.resource,m.visible) for m in Menu.objects.all()}
call_command('rebuild_business_navigation')
check('重复整理不修改接口/字段/数据范围授权', before=={m._meta.label:digest(m) for m in authorization_models})
check(str(len(RESOURCES))+'个资源业务行原样保留', data_before=={label:digest(apps.get_model(label)) for label in RESOURCES})
menus={m.component_name:m for m in Menu.objects.select_related('parent')}
check('原菜单ID、地址、资源及显示状态保留',all((menus[k].pk,menus[k].web_path,menus[k].resource,menus[k].visible)==v for k,v in menus_before.items()))


def domain(menu):
    while menu.parent is not None and menu.parent.component_name != 'erp':
        menu=menu.parent
    return menu.component_name


for label,spec in RESOURCES.items():
    check('归属 '+label,domain(menus[spec['entity']])==spec['module'],spec['module'])
check('试样业务仍保留打样模块',menus['erp_sampling_sourcedyedyelabmain'].parent.component_name=='erp_sampling_original')
for entity,title,entry in [('erp_sales_sourcedyeordermain','染整订单','OrderMain'),('erp_sales_sourcedyefixedprice','染整报价','Quote')]:
    response=client.get('/api/meta/'+entity+'/')
    check(title+'元数据和默认原入口',response.status_code==200 and response.data['data']['title']==title and response.data['data']['workflow']['default_entry']==entry)
    from system.catalog import ENTITIES
    view=view_for(entity,ENTITIES[entity])()
    from types import SimpleNamespace
    view.request=SimpleNamespace(method='GET',data={},query_params={},user=actor)
    check(title+'Excel默认入口与页面一致',view.excel_context()['业务入口']==entry)

rows_imported=field_comparisons=0
batch=None
for f in pack['files']:
    model=apps.get_model(f['model']);path=out/f['history_file']
    book=load_workbook(path,read_only=True,data_only=False)
    headers=[c.value for c in book['数据'][1]]
    check(f['title']+'原字段表头完整',headers==['序号',*[c['title'] for c in f['history_columns']]])
    cells=list(book['数据'].iter_rows(min_row=2))
    check(f['title']+'全部原行保留',len(cells)==max(1,len(f['rows'])) and all(c.data_type not in ['f','e'] for row in cells for c in row))
    book.close()
    url='/api/system/'+f['entity']+'/import/'+('? _foundation='+f['profile'] if f['profile'] else '')
    url=url.replace('? ','?')
    before_count=model.objects.count();actions=BusinessAction.objects.count();contexts=SourceDocumentContext.objects.count()
    response=client.post(url,{'file':SimpleUploadedFile(path.name,path.read_bytes())},format='multipart')
    check(f['title']+'实际接口试导',response.status_code==200 and response.data['data']['created']==len(f['rows']),str(response.data)[:700])
    check(f['title']+'不重演新业务事件',BusinessAction.objects.count()==actions and SourceDocumentContext.objects.count()==contexts)
    check(f['title']+'新增行数完整',model.objects.count()-before_count==len(f['rows']))
    if f['rows']:
        batch=LegacyImportBatch.objects.get(number=pack['batch'])
        saved={o.source_row:o for o in model.objects.filter(source_batch=batch,source_row__in=[r['row'] for r in f['rows']])}
        for raw in f['rows']:
            item=saved[raw['row']]
            for key,expected in original_values(model,raw).items():
                actual=getattr(item,key);actual=bytes(actual) if isinstance(actual,memoryview) else actual
                assert actual==expected,(f['title'],raw['row'],key,actual,expected)
                field_comparisons+=1
            assert item.approval_stage is None
        rows_imported+=len(saved)
    new=out/f['new_template']
    n=client.post(url,{'file':SimpleUploadedFile(new.name,new.read_bytes())},format='multipart')
    check(f['title']+'空白新增模板识别',n.status_code==200 and n.data['data']['created']==0,str(n.data)[:600])

detail=apps.get_model('erp_sales.SourceDyeFixedPriceDetail')
query=detail.objects.filter(source_batch=batch)
check('837个对客价及10个NULL原样保留',query.filter(old_cus_price__isnull=True).count()==10 and query.exclude(old_cus_price__isnull=True).exclude(old_cus_price=0).count()==837)
check('847个零总价与NULL客确价保留',query.filter(old_pricesum=0,old_price__isnull=True).count()==847)
check('报价主从自动关联完整',query.filter(source_parent_1__isnull=True).count()==0)
f=pack['files'][1];url='/api/system/'+f['entity']+'/import/';path=out/f['history_file'];n=detail.objects.count()
response=client.post(url,{'file':SimpleUploadedFile(path.name,path.read_bytes())},format='multipart')
check('同批重复提交明确失败且整批零写入',response.status_code==400 and detail.objects.count()==n)
check('888原行和23份xlsx全部覆盖',rows_imported==888 and len(list(out.rglob('*.xlsx')))==23)
report={'database':connection.settings_dict['NAME'],'port':5440,'checks':checks,'passed':len(checks),'rows_imported':rows_imported,
        'field_comparisons':field_comparisons,'source_batch':pack['batch'],'formal_imported':0,'preserved_authorizations':before,
        'scope':'独立库导入及开发验证；不代表正式历史对账或现场验收'}
(ROOT/'docs/erp-sales-production-verification-20261022.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print('ALL PASS',len(checks),'rows',rows_imported,'fields',field_comparisons)
